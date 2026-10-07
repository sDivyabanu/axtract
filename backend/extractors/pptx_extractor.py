"""PPTX extractor using python-pptx.

Preserves slide number, titles, body text, lists, tables, shapes, and images.
"""

from __future__ import annotations

from pathlib import Path
from typing import ClassVar

from pptx import Presentation
from pptx.util import Emu

from extractors.base import BaseExtractor, ExtractionResult
from models.document import BBox, BlockType, DocumentBlock
from models.errors import AppError, DocumentError


def _shape_bbox(shape, slide_w: int, slide_h: int) -> BBox | None:
    """Normalize shape position to 0.0–1.0 bbox."""
    if slide_w <= 0 or slide_h <= 0:
        return None
    try:
        return (
            round(shape.left / slide_w, 6),
            round(shape.top / slide_h, 6),
            round((shape.left + shape.width) / slide_w, 6),
            round((shape.top + shape.height) / slide_h, 6),
        )
    except Exception:
        return None


class PptxExtractor(BaseExtractor):
    """Extracts structured content from PPTX files."""

    name: ClassVar[str] = "python-pptx"
    supported_extensions: ClassVar[frozenset[str]] = frozenset({"pptx"})

    def extract(self, file_path: Path) -> ExtractionResult:
        try:
            prs = Presentation(str(file_path))
        except Exception as exc:
            raise AppError(
                "INVALID_FILE",
                f"Could not open PPTX file: {exc}",
                status_code=422,
            ) from exc

        slide_w = prs.slide_width or Emu(9144000)  # default 10"
        slide_h = prs.slide_height or Emu(6858000)  # default 7.5"

        blocks: list[DocumentBlock] = []
        errors: list[DocumentError] = []

        for slide_number, slide in enumerate(prs.slides, start=1):
            try:
                slide_blocks = self._extract_slide(
                    slide, slide_number, slide_w, slide_h
                )
                blocks.extend(slide_blocks)
            except Exception as exc:
                errors.append(
                    DocumentError(
                        code="SLIDE_EXTRACTION_FAILED",
                        message=str(exc),
                        page=slide_number,
                    )
                )

        return ExtractionResult(
            page_count=len(prs.slides),
            blocks=blocks,
            errors=errors,
        )

    def _extract_slide(
        self, slide, slide_number: int, slide_w: int, slide_h: int
    ) -> list[DocumentBlock]:
        blocks: list[DocumentBlock] = []
        block_counter = 0

        for shape in slide.shapes:
            # Title shapes
            if shape.has_text_frame:
                is_title = shape.shape_id == slide.shapes.title.shape_id if slide.shapes.title else False

                for para in shape.text_frame.paragraphs:
                    text = para.text.strip()
                    if not text:
                        continue

                    # Determine type
                    if is_title:
                        block_type = BlockType.HEADING
                    elif any(text.startswith(c) for c in ("•", "-", "–", "●", "○")):
                        block_type = BlockType.LIST
                    elif para.level > 0:
                        block_type = BlockType.LIST
                    else:
                        block_type = BlockType.PARAGRAPH

                    metadata: dict = {
                        "slide_number": slide_number,
                        "shape_name": shape.name,
                        "indent_level": para.level,
                    }

                    # Font info
                    if para.runs:
                        run = para.runs[0]
                        if run.font.size:
                            metadata["font_size"] = run.font.size.pt
                        metadata["is_bold"] = run.font.bold or False

                    blocks.append(
                        DocumentBlock(
                            id=f"s{slide_number}-b{block_counter}",
                            type=block_type,
                            content=text,
                            page=slide_number,
                            bbox=_shape_bbox(shape, slide_w, slide_h),
                            confidence=None,
                            extractor=self.name,
                            metadata=metadata,
                        )
                    )
                    block_counter += 1

            # Tables
            if shape.has_table:
                table = shape.table
                rows_data: list[list[str]] = []
                for row in table.rows:
                    cells = [cell.text.strip() for cell in row.cells]
                    rows_data.append(cells)

                if rows_data:
                    text_lines = [" | ".join(cells) for cells in rows_data]
                    blocks.append(
                        DocumentBlock(
                            id=f"s{slide_number}-b{block_counter}",
                            type=BlockType.TABLE,
                            content="\n".join(text_lines),
                            page=slide_number,
                            bbox=_shape_bbox(shape, slide_w, slide_h),
                            confidence=None,
                            extractor=self.name,
                            metadata={
                                "rows": rows_data,
                                "row_count": len(rows_data),
                                "col_count": len(rows_data[0]) if rows_data else 0,
                                "header": rows_data[0] if rows_data else [],
                                "slide_number": slide_number,
                            },
                        )
                    )
                    block_counter += 1

            # Images / pictures
            if shape.shape_type is not None and shape.shape_type == 13:  # MSO_SHAPE_TYPE.PICTURE
                blocks.append(
                    DocumentBlock(
                        id=f"s{slide_number}-b{block_counter}",
                        type=BlockType.FIGURE,
                        content=f"[image: {shape.name}]",
                        page=slide_number,
                        bbox=_shape_bbox(shape, slide_w, slide_h),
                        confidence=None,
                        extractor=self.name,
                        metadata={
                            "shape_name": shape.name,
                            "slide_number": slide_number,
                        },
                    )
                )
                block_counter += 1

        return blocks
