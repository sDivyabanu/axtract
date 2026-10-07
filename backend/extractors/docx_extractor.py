"""DOCX extractor using python-docx.

Preserves headings, paragraphs, lists, tables, and images as semantic blocks.
DOCX does not expose PDF-like page coordinates, so bbox is null.
"""

from __future__ import annotations

from pathlib import Path
from typing import ClassVar

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH

from extractors.base import BaseExtractor, ExtractionResult
from models.document import BlockType, DocumentBlock
from models.errors import AppError, DocumentError


class DocxExtractor(BaseExtractor):
    """Extracts structured content from DOCX files."""

    name: ClassVar[str] = "python-docx"
    supported_extensions: ClassVar[frozenset[str]] = frozenset({"docx"})

    def extract(self, file_path: Path) -> ExtractionResult:
        try:
            doc = docx.Document(str(file_path))
        except Exception as exc:
            raise AppError(
                "INVALID_FILE",
                f"Could not open DOCX file: {exc}",
                status_code=422,
            ) from exc

        blocks: list[DocumentBlock] = []
        errors: list[DocumentError] = []
        block_counter = 0
        page = 1  # DOCX doesn't expose real page numbers; use 1

        # Process body elements (paragraphs and tables in document order)
        for element in doc.element.body:
            try:
                tag = element.tag.split("}")[-1] if "}" in element.tag else element.tag

                if tag == "p":
                    para = None
                    for p in doc.paragraphs:
                        if p._element is element:
                            para = p
                            break
                    if para is None:
                        continue

                    block = self._paragraph_to_block(para, page, block_counter)
                    if block is not None:
                        blocks.append(block)
                        block_counter += 1

                elif tag == "tbl":
                    table = None
                    for t in doc.tables:
                        if t._element is element:
                            table = t
                            break
                    if table is None:
                        continue

                    block = self._table_to_block(table, page, block_counter)
                    if block is not None:
                        blocks.append(block)
                        block_counter += 1

            except Exception as exc:
                errors.append(
                    DocumentError(
                        code="ELEMENT_EXTRACTION_FAILED",
                        message=str(exc),
                        page=page,
                    )
                )

        # Extract images info
        for rel in doc.part.rels.values():
            if "image" in rel.reltype:
                blocks.append(
                    DocumentBlock(
                        id=f"p{page}-b{block_counter}",
                        type=BlockType.FIGURE,
                        content="[embedded image]",
                        page=page,
                        bbox=None,
                        confidence=None,
                        extractor=self.name,
                        metadata={"image_rel": rel.reltype},
                    )
                )
                block_counter += 1

        return ExtractionResult(
            page_count=1,  # DOCX doesn't reliably expose page count
            blocks=blocks,
            errors=errors,
        )

    def _paragraph_to_block(
        self, para, page: int, counter: int
    ) -> DocumentBlock | None:
        text = para.text.strip()
        if not text:
            return None

        style_name = (para.style.name or "").lower() if para.style else ""

        # Detect block type from style
        if "heading" in style_name or "title" in style_name:
            block_type = BlockType.HEADING
        elif "list" in style_name or "bullet" in style_name:
            block_type = BlockType.LIST
        else:
            block_type = BlockType.PARAGRAPH

        # Detect list by numbering or bullet characters
        if block_type == BlockType.PARAGRAPH:
            if text.startswith(("•", "-", "–", "—", "●", "○")):
                block_type = BlockType.LIST
            elif len(text) > 2 and text[0].isdigit() and text[1] in ".)" :
                block_type = BlockType.LIST

        metadata: dict = {"style": para.style.name if para.style else None}

        # Heading level
        if block_type == BlockType.HEADING and style_name:
            for i in range(1, 10):
                if str(i) in style_name:
                    metadata["heading_level"] = i
                    break

        # Font info from first run
        if para.runs:
            run = para.runs[0]
            if run.font.size:
                metadata["font_size"] = run.font.size.pt
            metadata["is_bold"] = run.bold or False
            metadata["is_italic"] = run.italic or False

        return DocumentBlock(
            id=f"p{page}-b{counter}",
            type=block_type,
            content=text,
            page=page,
            bbox=None,  # DOCX does not provide page coordinates
            confidence=None,  # Deterministic extraction
            extractor=self.name,
            metadata=metadata,
        )

    def _table_to_block(
        self, table, page: int, counter: int
    ) -> DocumentBlock | None:
        rows_data: list[list[str]] = []
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            rows_data.append(cells)

        if not rows_data:
            return None

        # Text representation
        text_lines = [" | ".join(cells) for cells in rows_data]
        content = "\n".join(text_lines)

        return DocumentBlock(
            id=f"p{page}-b{counter}",
            type=BlockType.TABLE,
            content=content,
            page=page,
            bbox=None,
            confidence=None,
            extractor=self.name,
            metadata={
                "rows": rows_data,
                "row_count": len(rows_data),
                "col_count": len(rows_data[0]) if rows_data else 0,
                "header": rows_data[0] if rows_data else [],
            },
        )
