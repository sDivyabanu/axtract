"""AXTRACT Verify layer 1: independent source inventories, built from real files."""

from __future__ import annotations

import re
from pathlib import Path

import pytest

from tests.verify_fixtures import extract, make_docx, make_pdf, make_pptx, make_xlsx, png_file
from verify.inventory import (
    Contract, InventoryOptions, SourceInventory, SourceObjectType as SO, build_inventory,
)
from verify.models import EvidenceKind, UnitType

DET, HEUR, NA = EvidenceKind.DETERMINISTIC, EvidenceKind.HEURISTIC, EvidenceKind.UNAVAILABLE


@pytest.fixture(scope="module")
def tmp(tmp_path_factory):
    return tmp_path_factory.mktemp("inventory")


@pytest.fixture(scope="module")
def xlsx_inv(tmp):
    return build_inventory(make_xlsx(tmp / "book.xlsx"), "xlsx")


@pytest.fixture(scope="module")
def pptx_inv(tmp):
    return build_inventory(make_pptx(tmp / "deck.pptx"), "pptx")


@pytest.fixture(scope="module")
def docx_inv(tmp):
    return build_inventory(make_docx(tmp / "doc.docx"), "docx")


@pytest.fixture(scope="module")
def pdf_inv(tmp):
    return build_inventory(make_pdf(tmp / "doc.pdf", pages=3, scanned_page=3), "pdf")


def objs(inv: SourceInventory, type_: SO, unit_index: int | None = None):
    return [o for o in inv.objects(type_) if unit_index is None or o.unit.index == unit_index]


class TestXlsxInventory:
    def test_sheets_in_workbook_order_with_names_and_visibility(self, xlsx_inv):
        assert xlsx_inv.error is None and xlsx_inv.properties["sheet_count"] == 4
        assert [(u.unit.index, u.unit.label, u.properties["state"]) for u in xlsx_inv.units] == [
            (1, "Sales", "visible"), (2, "Costs", "visible"), (3, "Hidden notes", "hidden"), (4, "Empty", "visible")]
        assert all(u.unit.type == UnitType.SHEET for u in xlsx_inv.units)

    def test_populated_cells_coordinates_and_cached_values_from_raw_xml(self, xlsx_inv):
        region = objs(xlsx_inv, SO.CELL_REGION, 1)[0]
        cells = region.metadata["cells"]
        assert region.locator.cell_range == "A1:F4" and region.metadata["non_empty_cells"] == 14
        assert cells["A1"] == "Region" and cells["B2"] == "10" and cells["A3"] == "South" and cells["F1"] == "Merged heading"

    def test_formula_presence_and_coordinates_are_recorded_but_outside_the_contract(self, xlsx_inv):
        f = objs(xlsx_inv, SO.FORMULA, 1)[0]
        assert f.metadata["coordinates"] == ["D2"] and f.metadata["without_cached_value"] == 1
        assert f.contract == Contract.NOT_IN_CONTRACT and "cached values" in f.contract_note
        assert "D2" not in objs(xlsx_inv, SO.CELL_REGION, 1)[0].metadata["cells"]  # nothing to lose without a cached value

    def test_merged_ranges(self, xlsx_inv):
        assert [o.locator.cell_range for o in objs(xlsx_inv, SO.MERGED_RANGE)] == ["F1:G2"]

    def test_excel_table_definition(self, xlsx_inv):
        t = objs(xlsx_inv, SO.TABLE)[0]
        assert t.locator.cell_range == "A1:C4" and t.text == "SalesTbl"

    def test_chart_and_picture_with_anchors(self, xlsx_inv):
        c, p = objs(xlsx_inv, SO.CHART)[0], objs(xlsx_inv, SO.PICTURE)[0]
        assert c.metadata["chart_types"] == ["barChart"] and c.text == "Sales chart" and c.locator.cell_range == "J2"
        assert p.locator.cell_range == "J20" and p.metadata["media"].startswith("xl/media/")

    def test_hidden_rows_hidden_sheet_and_empty_sheet(self, xlsx_inv):
        assert xlsx_inv.unit(1).properties["hidden_rows"] == [4]
        assert objs(xlsx_inv, SO.CELL_REGION, 3) and not objs(xlsx_inv, SO.CELL_REGION, 4)

    def test_hidden_columns(self, tmp):
        import openpyxl

        wb = openpyxl.Workbook()
        wb.active["A1"], wb.active["B1"] = "x", "y"
        wb.active.column_dimensions["B"].hidden = True
        wb.save(tmp / "hidecol.xlsx")
        assert build_inventory(tmp / "hidecol.xlsx", "xlsx").unit(1).properties["hidden_columns"] == [(2, 2)]

    def test_every_object_states_how_it_was_observed(self, xlsx_inv):
        for o in xlsx_inv.objects():
            assert o.kind == DET and o.engine.startswith("ooxml:"), o

    def test_workbooks_without_drawings_have_none(self, tmp):
        inv = build_inventory(make_xlsx(tmp / "plain.xlsx", chart=False, picture=False, table=False, merge=False, formulas=False), "xlsx")
        assert not list(inv.objects(SO.CHART)) and not list(inv.objects(SO.PICTURE)) and not list(inv.objects(SO.TABLE))

    def test_shared_and_inline_strings_and_rich_text(self, tmp):
        import openpyxl
        from openpyxl.cell.rich_text import CellRichText, TextBlock
        from openpyxl.cell.text import InlineFont

        wb = openpyxl.Workbook()
        wb.active["A1"] = CellRichText(TextBlock(InlineFont(b=True), "Bold "), "plain")
        wb.active["A2"] = True
        wb.save(tmp / "rich.xlsx")
        cells = objs(build_inventory(tmp / "rich.xlsx", "xlsx"), SO.CELL_REGION)[0].metadata["cells"]
        assert cells["A1"] == "Bold plain" and cells["A2"] == "1"

    def test_the_cell_cap_is_honoured_and_reported(self, tmp):
        inv = build_inventory(tmp / "book.xlsx", "xlsx", InventoryOptions(max_cells=5))
        assert inv.truncated and inv.unit(1).properties["truncated"] is True

    def test_dates_are_reported_as_raw_serials_and_the_limitation_says_so(self, tmp):
        import datetime, openpyxl

        wb = openpyxl.Workbook()
        wb.active["A1"] = datetime.date(2024, 1, 2)
        wb.save(tmp / "dates.xlsx")
        inv = build_inventory(tmp / "dates.xlsx", "xlsx")
        assert objs(inv, SO.CELL_REGION)[0].metadata["cells"]["A1"] == "45293"
        assert any("serial numbers" in s for s in inv.limitations)


class TestPptxInventory:
    def test_slide_count_order_and_size(self, pptx_inv):
        assert pptx_inv.error is None and pptx_inv.properties["slide_count"] == 2
        assert [u.unit.type for u in pptx_inv.units] == [UnitType.SLIDE] * 2 and pptx_inv.properties["slide_width_emu"] > 0

    def test_slide_one_objects(self, pptx_inv):
        by = {o.locator.shape_name: o for o in pptx_inv.unit(1).objects}
        assert by["Summary box"].text == "Revenue grew strongly across all regions this quarter" and by["Summary box"].type == SO.TEXT
        t = by["Results table"]
        assert (t.type, t.metadata["row_count"], t.metadata["col_count"]) == (SO.TABLE, 3, 3) and t.metadata["rows"][1][2] == "r1c2"
        assert by["Team photo"].type == SO.PICTURE
        title = next(o for o in pptx_inv.unit(1).objects if o.type == SO.HEADING)
        assert title.text == "Quarterly Review"

    def test_chart_group_and_grouped_text_on_slide_two(self, pptx_inv):
        by = {o.locator.shape_name: o for o in pptx_inv.unit(2).objects}
        assert by["Revenue chart"].type == SO.CHART and by["Revenue chart"].metadata["series_count"] == 1
        assert by["Callout group"].type == SO.GROUP and by["Callout group"].contract == Contract.NOT_IN_CONTRACT
        assert by["Group note"].text == "Text that lives inside a group shape"

    def test_geometry_is_exact_and_matches_the_extractors_own_boxes(self, pptx_inv, tmp):
        resp = extract(tmp / "deck.pptx")
        for name in ("Summary box", "Results table", "Team photo", "Group note"):
            o = next(x for x in pptx_inv.objects() if x.locator.shape_name == name)
            blocks = [b for b in resp.blocks if b.metadata.get("shape_name") == name or (name == "Results table" and b.type.value == "table")]
            assert o.locator.bbox and blocks and blocks[0].bbox
            assert all(abs(a - b) < 1e-3 for a, b in zip(o.locator.bbox, blocks[0].bbox)), name

    def test_group_children_are_transformed_into_slide_space(self, pptx_inv):
        by = {o.locator.shape_name: o for o in pptx_inv.unit(2).objects}
        inner, group = by["Group note"].locator.bbox, by["Callout group"].locator.bbox
        assert inner and 0 <= inner[0] < inner[2] <= 1 and 0 <= inner[1] < inner[3] <= 1
        assert group is None or (group[0] <= inner[0] + 1e-6 and inner[2] <= group[2] + 1e-6)

    def test_layout_inherited_placeholders_have_no_invented_bbox(self, pptx_inv):
        assert next(o for o in pptx_inv.unit(1).objects if o.type == SO.HEADING).locator.bbox is None

    def test_speaker_notes_are_recorded_but_outside_the_contract(self, pptx_inv):
        n = next(o for o in pptx_inv.unit(1).objects if o.metadata.get("role") == "speaker_notes")
        assert n.text == "Remember to mention the churn numbers" and n.contract == Contract.NOT_IN_CONTRACT

    def test_every_object_states_how_it_was_observed(self, pptx_inv):
        for o in pptx_inv.objects():
            assert o.kind == DET and o.engine.startswith("ooxml:"), o

    def test_hidden_slides_are_flagged(self, tmp):
        import pptx

        p = make_pptx(tmp / "hid.pptx", picture=False, chart=False, table=False, group=False, notes=False)
        prs = pptx.Presentation(p)
        prs.slides[1]._element.set("show", "0")
        prs.save(p)
        inv = build_inventory(p, "pptx")
        assert [u.properties["hidden"] for u in inv.units] == [False, True]


class TestDocxInventory:
    def test_objects_in_document_order_with_logical_positions(self, docx_inv):
        o = docx_inv.units[0].objects
        assert [x.type for x in o if x.contract == Contract.EXPECTED] == [
            SO.HEADING, SO.TEXT, SO.HEADING, SO.TEXT, SO.LIST_ITEM, SO.LIST_ITEM, SO.TABLE, SO.PICTURE, SO.EQUATION, SO.TEXT]
        body = [x for x in o if x.contract == Contract.EXPECTED]
        positions = [x.locator.position for x in body]
        assert positions == sorted(positions) and all(x.locator.path.startswith("body[") for x in body)
        assert all(x.locator.position is None for x in o if x.contract == Contract.NOT_IN_CONTRACT)  # header: not in the body

    def test_text_headings_and_lists(self, docx_inv):
        o = [x for x in docx_inv.units[0].objects if x.type in (SO.HEADING, SO.LIST_ITEM)]
        assert [(x.type, x.text) for x in o] == [
            (SO.HEADING, "Annual Report"), (SO.HEADING, "Financial Results"),
            (SO.LIST_ITEM, "First list entry about growth"), (SO.LIST_ITEM, "Second list entry about costs")]

    def test_table_dimensions_and_cell_text(self, docx_inv):
        t = next(docx_inv.objects(SO.TABLE))
        assert (t.metadata["row_count"], t.metadata["col_count"]) == (3, 3) and t.metadata["rows"][2][1] == "cell21"

    def test_equation_is_detected_from_omml(self, docx_inv):
        e = next(docx_inv.objects(SO.EQUATION))
        assert e.text == "E=mc2" and e.metadata["display"] is True

    def test_picture_is_detected(self, docx_inv):
        assert len(list(docx_inv.objects(SO.PICTURE))) == 1

    def test_headers_are_recorded_but_outside_the_contract(self, docx_inv):
        h = next(o for o in docx_inv.objects() if o.metadata.get("role") == "header")
        assert h.text == "Confidential draft" and h.contract == Contract.NOT_IN_CONTRACT

    def test_no_page_or_bbox_is_ever_invented(self, docx_inv):
        for o in docx_inv.objects():
            assert o.locator.page is None and o.locator.bbox is None, o
        sig = next(s for s in docx_inv.signals if s.name == "rendered_page_geometry")
        assert sig.kind == NA and not sig.available

    def test_sections_are_counted(self, docx_inv):
        assert docx_inv.properties["sections"] == 1 and docx_inv.units[0].unit.type == UnitType.SECTION

    def test_every_object_states_how_it_was_observed(self, docx_inv):
        for o in docx_inv.objects():
            assert o.kind == DET and o.engine.startswith("ooxml:"), o

    def test_tracked_deletions_are_not_counted_as_text(self, tmp):
        import docx
        from docx.oxml import parse_xml

        d = docx.Document()
        p = d.add_paragraph("kept")
        p._p.append(parse_xml('<w:del xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" w:id="1" w:author="a">'
                              '<w:r><w:delText>deleted</w:delText></w:r></w:del>'))
        d.save(tmp / "del.docx")
        assert next(build_inventory(tmp / "del.docx", "docx").objects(SO.TEXT)).text == "kept"


class TestPdfInventory:
    def test_pages_text_and_classes(self, pdf_inv):
        assert pdf_inv.error is None and pdf_inv.properties["page_count"] == 3
        assert [(u.unit.index, u.properties["page_class"]) for u in pdf_inv.units] == [(1, "digital"), (2, "digital"), (3, "scanned")]
        page = next(o for o in objs(pdf_inv, SO.TEXT, 1) if o.metadata["scope"] == "page")
        assert "Section 1 Overview" in page.text and (page.kind, page.engine) == (DET, "pypdfium2")

    def test_text_regions_are_heuristic_groupings_of_exact_words(self, pdf_inv):
        regions = [o for o in objs(pdf_inv, SO.TEXT, 1) if o.metadata["scope"] == "region"]
        assert regions and all((o.kind, o.engine) == (HEUR, "pdfplumber") and o.locator.bbox for o in regions)
        assert all(o.metadata["text_kind"] == "deterministic" and o.metadata["grouping_kind"] == "heuristic" for o in regions)

    def test_table_candidate_is_labelled_a_candidate(self, pdf_inv):
        t = objs(pdf_inv, SO.TABLE, 1)[0]
        assert (t.kind, t.engine) == (HEUR, "pdfplumber") and (t.metadata["row_count"], t.metadata["col_count"]) == (4, 3)
        assert "North" in t.text and "not ground truth" in t.metadata["limitation"] and t.locator.bbox

    def test_image_objects_have_deterministic_bounds(self, pdf_inv):
        imgs = objs(pdf_inv, SO.PICTURE, 1)
        assert imgs and all((o.kind, o.engine) == (DET, "pypdfium2") and o.locator.bbox for o in imgs)
        assert all(0 < o.metadata["area_fraction"] <= 1 for o in imgs)

    def test_a_scanned_page_has_no_independent_text_reader(self, pdf_inv):
        u = pdf_inv.unit(3)
        assert u.independent_text_available is False and "OCR" in u.independence_note
        assert not [o for o in u.objects if o.type == SO.TEXT]
        assert next(o for o in u.objects if o.type == SO.PICTURE).metadata["full_page_background"] is True

    def test_page_class_is_labelled_heuristic(self, pdf_inv):
        assert all(u.properties["page_class_kind"] == "heuristic" for u in pdf_inv.units)
        assert next(s for s in pdf_inv.signals if s.name == "page_class_digital_scanned_mixed").kind == HEUR

    def test_unavailable_signals_are_declared_unavailable(self, pdf_inv):
        for name in ("scanned_page_text", "charts_figures_equations"):
            s = next(s for s in pdf_inv.signals if s.name == name)
            assert s.kind == NA and s.available is False

    def test_the_detailed_page_limit_truncates_honestly(self, tmp):
        inv = build_inventory(tmp / "doc.pdf", "pdf", InventoryOptions(max_pdf_pages_detailed=1))
        assert inv.truncated and "native text only" in inv.unit(2).properties["detail"]
        assert not [o for o in inv.unit(2).objects if o.metadata.get("scope") == "region"]
        assert [o for o in inv.unit(2).objects if o.metadata.get("scope") == "page"]  # native text is still read

    def test_corrupt_and_non_pdf_files_report_an_error_instead_of_raising(self, tmp):
        (tmp / "bad.pdf").write_bytes(b"not a pdf at all")
        inv = build_inventory(tmp / "bad.pdf", "pdf")
        assert inv.error and inv.units == []

    def test_does_not_use_the_extractors_pdf_engine(self):
        src = Path(__file__).resolve().parents[1].joinpath("verify", "inventory", "pdf.py").read_text(encoding="utf-8")
        assert not re.search(r"^\s*(import|from)\s+(pymupdf|fitz)\b", src, re.M)


class TestImageInventory:
    def test_dimensions_and_properties_only(self, tmp):
        p = png_file(tmp / "pic.png", (300, 200))
        inv = build_inventory(p, "png")
        assert inv.error is None and inv.properties["width_px"] == 300 and inv.properties["height_px"] == 200
        assert inv.properties["format"] == "PNG" and inv.units[0].unit.type == UnitType.IMAGE
        assert [o.type for o in inv.objects()] == [SO.PICTURE]

    def test_independent_text_evidence_is_declared_unavailable(self, tmp):
        inv = build_inventory(png_file(tmp / "pic2.png"), "png")
        u = inv.units[0]
        assert u.independent_text_available is False and "same engine" in u.independence_note or "engine AXTRACT used" in u.independence_note
        assert next(s for s in inv.signals if s.name == "text_regions").kind == NA

    def test_jpeg(self, tmp):
        from PIL import Image

        Image.new("RGB", (50, 40)).save(tmp / "p.jpg", "JPEG")
        assert build_inventory(tmp / "p.jpg", "jpg").properties["format"] == "JPEG"

    def test_a_broken_image_reports_an_error(self, tmp):
        (tmp / "broken.png").write_bytes(b"\x89PNG nope")
        assert build_inventory(tmp / "broken.png", "png").error


class TestEverySignalStatesItsEvidence:
    @pytest.mark.parametrize("fixture", ["xlsx_inv", "pptx_inv", "docx_inv", "pdf_inv"])
    def test_signals_have_a_kind_an_engine_and_a_note_when_weak(self, fixture, request):
        inv = request.getfixturevalue(fixture)
        assert inv.signals and inv.limitations
        for s in inv.signals:
            assert isinstance(s.kind, EvidenceKind) and s.engine, s
            if s.kind in (HEUR, NA):
                assert s.note, f"{s.name}: weak evidence must explain itself"
            assert s.available == (s.kind != NA)

    def test_image_signals(self, tmp):
        inv = build_inventory(png_file(tmp / "s.png"), "png")
        assert {s.name: s.kind for s in inv.signals}["format_dimensions_mode"] == DET


class TestRobustness:
    def test_unknown_format(self, tmp):
        (tmp / "x.txt").write_text("hi")
        assert build_inventory(tmp / "x.txt", "txt").error

    @pytest.mark.parametrize("ext", ["xlsx", "pptx", "docx"])
    def test_non_zip_office_files(self, tmp, ext):
        p = tmp / f"fake.{ext}"
        p.write_bytes(b"this is not a zip")
        inv = build_inventory(p, ext)
        assert inv.error and not inv.units

    def test_a_zip_missing_the_main_part(self, tmp):
        import zipfile

        p = tmp / "empty.xlsx"
        with zipfile.ZipFile(p, "w") as z:
            z.writestr("hello.txt", "x")
        assert build_inventory(p, "xlsx").error

    def test_inventories_are_deterministic(self, tmp):
        a = build_inventory(tmp / "deck.pptx", "pptx")
        b = build_inventory(tmp / "deck.pptx", "pptx")
        assert [o.model_dump() for o in a.objects()] == [o.model_dump() for o in b.objects()]

    def test_timing_is_recorded(self, xlsx_inv, pptx_inv, docx_inv, pdf_inv):
        assert all(i.timing_ms > 0 for i in (xlsx_inv, pptx_inv, docx_inv, pdf_inv))


class TestIndependenceFromTheExtractors:
    """The inventory must not read the document through the libraries AXTRACT's extractors use."""

    BANNED = {"xlsx.py": ["openpyxl"], "pptx.py": ["pptx"], "docx.py": ["docx"], "pdf.py": ["pymupdf", "fitz"],
              "ooxml.py": ["openpyxl", "pptx", "docx"]}

    @pytest.mark.parametrize("module,banned", BANNED.items())
    def test_no_extractor_libraries_are_imported(self, module, banned):
        src = Path(__file__).resolve().parents[1].joinpath("verify", "inventory", module).read_text(encoding="utf-8")
        for lib in banned:
            assert not re.search(rf"^\s*(import|from)\s+{lib}\b", src, re.M), f"{module} imports {lib}"

    def test_the_inventory_package_never_imports_the_extractors(self):
        folder = Path(__file__).resolve().parents[1] / "verify" / "inventory"
        for f in folder.glob("*.py"):
            assert not re.search(r"^\s*(import|from)\s+extractors\b", f.read_text(encoding="utf-8"), re.M), f.name
