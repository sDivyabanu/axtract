"""Generators for REAL files used by the inventory / completeness tests, plus helpers to extract them.

Nothing is mocked: every document is a genuine file, parsed by AXTRACT's real pipeline, so the
tests exercise the same path as production.
"""

from __future__ import annotations

import io
from pathlib import Path

from PIL import Image, ImageDraw
from starlette.datastructures import UploadFile

from models.document import DocumentResponse
from services.parse_service import parse_upload


def png_file(path: Path, size=(120, 80), colour="steelblue") -> Path:
    Image.new("RGB", size, colour).save(path, "PNG")
    return path


def extract(path: Path) -> DocumentResponse:
    """Run AXTRACT's real pipeline on a file."""
    return parse_upload(UploadFile(file=io.BytesIO(path.read_bytes()), filename=path.name))


def drop(resp: DocumentResponse, pred) -> DocumentResponse:
    """A copy of the extraction with every block matching `pred` removed: a simulated silent loss."""
    out = resp.model_copy(deep=True)
    out.blocks = [b for b in out.blocks if not pred(b)]
    return out


# --------------------------------------------------------------------------------------- XLSX


def make_xlsx(path: Path, *, chart=True, picture=True, table=True, merge=True, hidden_sheet=True, formulas=True,
              empty_sheet=True, hidden_row=True) -> Path:
    import openpyxl
    from openpyxl.chart import BarChart, Reference
    from openpyxl.drawing.image import Image as XImage
    from openpyxl.worksheet.table import Table

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sales"
    ws.append(["Region", "Q1", "Q2"])
    ws.append(["North", 10, 20])
    ws.append(["South", 30, 40])
    ws.append(["West", 50, 60])
    if formulas:
        ws["D1"], ws["D2"] = "Total", "=B2+C2"
    if merge:
        ws["F1"] = "Merged heading"
        ws.merge_cells("F1:G2")
    if hidden_row:
        ws.row_dimensions[4].hidden = True
    if table:
        ws.add_table(Table(displayName="SalesTbl", ref="A1:C4"))
    if chart:
        ch = BarChart()
        ch.title = "Sales chart"
        ch.add_data(Reference(ws, min_col=2, min_row=1, max_row=4, max_col=3), titles_from_data=True)
        ws.add_chart(ch, "J2")
    if picture:
        png = png_file(path.with_suffix(".logo.png"), (40, 30), "red")
        ws.add_image(XImage(str(png)), "J20")
    costs = wb.create_sheet("Costs")
    costs.append(["Item", "Amount"])
    costs.append(["Rent", 1200])
    costs.append(["Power", 340])
    if hidden_sheet:
        h = wb.create_sheet("Hidden notes")
        h["A1"], h["B1"] = "internal", "memo"
        h.sheet_state = "hidden"
    if empty_sheet:
        wb.create_sheet("Empty")
    wb.save(path)
    return path


# --------------------------------------------------------------------------------------- PPTX


def make_pptx(path: Path, *, picture=True, chart=True, table=True, group=True, notes=True) -> Path:
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.util import Inches

    prs = Presentation()
    s1 = prs.slides.add_slide(prs.slide_layouts[5])  # title only
    s1.shapes.title.text = "Quarterly Review"
    tb = s1.shapes.add_textbox(Inches(1), Inches(1.6), Inches(6), Inches(1))
    tb.name = "Summary box"
    tb.text_frame.text = "Revenue grew strongly across all regions this quarter"
    if table:
        t = s1.shapes.add_table(3, 3, Inches(1), Inches(3), Inches(6), Inches(1.5))
        t.name = "Results table"
        for r in range(3):
            for c in range(3):
                t.table.cell(r, c).text = f"r{r}c{c}"
    if picture:
        png = png_file(path.with_suffix(".pic.png"), (120, 80), "green")
        s1.shapes.add_picture(str(png), Inches(7.5), Inches(1.5), Inches(1.5), Inches(1)).name = "Team photo"
    if notes:
        s1.notes_slide.notes_text_frame.text = "Remember to mention the churn numbers"
    s2 = prs.slides.add_slide(prs.slide_layouts[5])
    s2.shapes.title.text = "Trends"
    if chart:
        cd = CategoryChartData()
        cd.categories = ["Q1", "Q2", "Q3"]
        cd.add_series("Revenue", (10, 20, 30))
        gf = s2.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(2), Inches(5), Inches(3), cd)
        gf.name = "Revenue chart"
    if group:
        g = s2.shapes.add_group_shape()
        g.name = "Callout group"
        inner = g.shapes.add_textbox(Inches(6), Inches(2), Inches(3), Inches(1))
        inner.name = "Group note"
        inner.text_frame.text = "Text that lives inside a group shape"
    prs.save(path)
    return path


# --------------------------------------------------------------------------------------- DOCX


_OMML = ('<m:oMathPara xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math">'
         '<m:oMath><m:r><m:t>E=mc2</m:t></m:r></m:oMath></m:oMathPara>')


def make_docx(path: Path, *, picture=True, table=True, equation=True, header=True, lists=True) -> Path:
    import docx
    from docx.oxml import parse_xml

    d = docx.Document()
    d.add_heading("Annual Report", level=1)
    d.add_paragraph("The company expanded into three new markets during the year.")
    d.add_heading("Financial Results", level=2)
    d.add_paragraph("Revenue rose strongly while operating costs remained flat.")
    if lists:
        d.add_paragraph("First list entry about growth", style="List Bullet")
        d.add_paragraph("Second list entry about costs", style="List Bullet")
    if table:
        t = d.add_table(rows=3, cols=3)
        for r in range(3):
            for c in range(3):
                t.cell(r, c).text = f"cell{r}{c}"
    if picture:
        d.add_picture(str(png_file(path.with_suffix(".pic.png"), (100, 60), "orange")))
    if equation:
        p = d.add_paragraph()
        p._p.append(parse_xml(_OMML))
    d.add_paragraph("Closing remarks summarise the outlook for the coming year.")
    if header:
        d.sections[0].header.paragraphs[0].text = "Confidential draft"
    d.save(path)
    return path


# --------------------------------------------------------------------------------------- PDF


PARAS = [
    "Revenue increased strongly across every region during the fiscal year.",
    "Operating costs remained broadly flat despite expansion into new markets.",
    "Management expects continued growth driven by the new product lines.",
]


def make_pdf(path: Path, *, pages=2, table=True, image=True, scanned_page: int | None = None) -> Path:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas
    from reportlab.platypus import Table, TableStyle

    c = canvas.Canvas(str(path), pagesize=A4)
    w, h = A4
    for pg in range(1, pages + 1):
        if scanned_page == pg:
            img = Image.new("RGB", (1240, 1754), "white")
            d = ImageDraw.Draw(img)
            for i, line in enumerate(PARAS):
                d.text((100, 200 + i * 80), line, fill="black")
            tmp = path.with_suffix(f".scan{pg}.png")
            img.save(tmp)
            c.drawImage(str(tmp), 0, 0, w, h)
        else:
            c.setFont("Helvetica-Bold", 16)
            c.drawString(60, h - 60, f"Section {pg} Overview")
            c.setFont("Helvetica", 11)
            for i, para in enumerate(PARAS):
                c.drawString(60, h - 110 - i * 40, f"{para} (page {pg})")
            if table and pg == 1:
                t = Table([["Region", "Q1", "Q2"], ["North", "10", "20"], ["South", "30", "40"], ["West", "50", "60"]], colWidths=[120, 80, 80])
                t.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.8, colors.black)]))
                t.wrapOn(c, w, h)
                t.drawOn(c, 60, h - 420)
            if image and pg == 1:
                png = png_file(path.with_suffix(".img.png"), (200, 140), "purple")
                c.drawImage(str(png), 60, h - 640, 240, 170)
        c.showPage()
    c.save()
    return path
