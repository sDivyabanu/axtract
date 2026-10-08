"""Live processing progress: real events, no leaks, nothing changed for the non-streaming endpoints."""

from __future__ import annotations

import asyncio
import json

import pytest
from fastapi.testclient import TestClient

from services import progress
from services.progress import STAGE_IDS, Tracker, _clean_detail
from services.progress_stream import event_stream
from tests.test_documents_api import FakeDb, env  # noqa: F401  (env: fake DB + storage + real crypto)
from tests.verify_fixtures import make_docx, make_pdf

TERMINAL = {"completed", "warning", "failed", "skipped"}
MARKER = "Zyxwvut-unique-marker-4471"
WHITELIST = {"findings", "issues", "blocks", "pages", "errors", "by_severity", "codes", "reason", "error_type", "status"}


@pytest.fixture(scope="module")
def tmp(tmp_path_factory):
    return tmp_path_factory.mktemp("progress")


@pytest.fixture(scope="module")
def docx(tmp):
    return make_docx(tmp / "secret-name-report.docx")


@pytest.fixture
def client():
    import main

    main.app.dependency_overrides.clear()
    with TestClient(main.app) as c:
        yield c


def parse_sse(raw: str) -> list[tuple[str, object]]:
    out = []
    for block in raw.split("\n\n"):
        event, data = None, []
        for line in block.splitlines():
            if line.startswith("event: "):
                event = line[7:]
            elif line.startswith("data: "):
                data.append(line[6:])
        if event:
            out.append((event, json.loads("".join(data))))
    return out


def stream(client, path, url="/api/parse/stream"):
    with open(path, "rb") as fh:
        r = client.post(url, files={"file": (path.name, fh, "application/octet-stream")})
    assert r.headers["content-type"].startswith("text/event-stream")
    return r.text, parse_sse(r.text)


def stage_events(events):
    return [d for e, d in events if e == "stage"]


# ================================================================== the stream


class TestStream:
    def test_a_real_run_emits_every_stage_in_order_and_then_the_result(self, client, docx):
        _, events = stream(client, docx)
        names = [e for e, _ in events]
        assert names[0] == "start" and names[-1] == "result"
        assert [s["id"] for s in events[0][1]["stages"]] == STAGE_IDS
        stages = stage_events(events)
        assert [s["seq"] for s in stages] == sorted(s["seq"] for s in stages)
        first_running = [next(i for i, s in enumerate(stages) if s["id"] == sid and s["state"] == "running") for sid in STAGE_IDS]
        assert first_running == sorted(first_running), "stages start in the order they really run"
        final = {s["id"]: s["state"] for s in stages}
        assert set(final) == set(STAGE_IDS) and all(v in TERMINAL for v in final.values())
        for sid in STAGE_IDS:  # a stage never reports an outcome before it started
            assert [s["state"] for s in stages if s["id"] == sid][0] == "running"

    def test_the_streamed_result_matches_the_plain_endpoint(self, client, docx):
        _, events = stream(client, docx)
        streamed = events[-1][1]
        with open(docx, "rb") as fh:
            plain = client.post("/api/parse", files={"file": (docx.name, fh, "application/octet-stream")}).json()
        assert [b["content"] for b in streamed["blocks"]] == [b["content"] for b in plain["blocks"]]
        assert streamed["validation"]["status"] == plain["validation"]["status"]
        assert streamed["security_findings"] == plain["security_findings"]

    def test_the_verdict_event_agrees_with_the_report(self, client, docx):
        _, events = stream(client, docx)
        verdict = [s for s in stage_events(events) if s["id"] == "verdict"][-1]
        assert verdict["detail"]["status"] == events[-1][1]["validation"]["status"]

    def test_issue_counts_in_stage_events_match_the_report(self, client, tmp):
        pdf = make_pdf(tmp / "a.pdf", pages=2)
        _, events = stream(client, pdf)
        report = events[-1][1]["validation"]
        last = {}
        for s in stage_events(events):
            if s["id"] in ("checks", "content", "structure") and "detail" in s:
                last[s["id"]] = s["detail"].get("issues", 0)
        assert sum(last.values()) == len(report["issues"])

    def test_the_plain_endpoint_is_unchanged(self, client, docx):
        with open(docx, "rb") as fh:
            r = client.post("/api/parse", files={"file": (docx.name, fh, "application/octet-stream")})
        assert r.status_code == 200 and r.headers["content-type"].startswith("application/json")
        assert r.json()["validation"]["status"]

    def test_validation_disabled_marks_those_stages_skipped_and_says_why(self, client, docx, monkeypatch):
        monkeypatch.setenv("AXTRACT_VERIFY", "0")
        _, events = stream(client, docx)
        assert events[0][1]["verify"] is False
        final = {s["id"]: s for s in stage_events(events)}
        for sid in ("inventory", "checks", "content", "structure", "verdict"):
            assert final[sid]["state"] == "skipped" and final[sid]["detail"]["reason"] == "disabled"
        assert final["security"]["state"] != "skipped" and final["extraction"]["state"] != "skipped"
        assert events[-1][1]["validation"] is None and events[-1][1]["blocks"]

    def test_a_verify_crash_is_reported_and_the_extraction_survives(self, client, docx, monkeypatch):
        import verify.engine as eng

        monkeypatch.setattr(eng, "run_verification", lambda *a, **k: (_ for _ in ()).throw(RuntimeError("boom")))
        _, events = stream(client, docx)
        final = {s["id"]: s["state"] for s in stage_events(events)}
        assert final["verdict"] == "failed" and events[-1][0] == "result" and events[-1][1]["blocks"]

    def test_a_failing_layer_is_a_failed_stage_and_its_dependants_are_skipped(self, client, docx, monkeypatch):
        import verify.engine as eng

        monkeypatch.setattr(eng, "build_inventory", lambda *a, **k: (_ for _ in ()).throw(ValueError("no")))
        _, events = stream(client, docx)
        final = {s["id"]: s for s in stage_events(events)}
        assert final["inventory"]["state"] == "failed" and final["inventory"]["detail"]["error_type"] == "ValueError"
        assert final["content"]["state"] == "skipped" and final["content"]["detail"]["reason"] == "upstream_failed"
        assert events[-1][0] == "result" and events[-1][1]["blocks"]

    def test_errors_end_the_stream_with_an_error_event_and_no_result(self, client, tmp):
        bad = tmp / "x.txt"
        bad.write_text("hello")
        _, events = stream(client, bad)
        assert events[-1][0] == "error" and events[-1][1]["code"] == "UNSUPPORTED_FORMAT"
        assert "result" not in [e for e, _ in events]

    def test_a_corrupt_pdf_is_a_422_error_event(self, client, tmp):
        bad = tmp / "bad.pdf"
        bad.write_bytes(b"%PDF-1.7\n" + b"garbage" * 50)
        _, events = stream(client, bad)
        assert events[-1][0] == "error" and events[-1][1]["status_code"] == 422

    def test_the_upload_is_removed_after_a_streamed_parse(self, client, docx, monkeypatch):
        import services.parse_service as ps

        seen = []
        real = ps.remove_temp_file
        monkeypatch.setattr(ps, "remove_temp_file", lambda p: (seen.append(p), real(p))[1])
        stream(client, docx)
        assert seen and not any(p and p.exists() for p in seen)


# ================================================================== no leaks


class TestNoLeaks:
    def test_stage_events_carry_no_document_text_names_or_paths(self, client, tmp):
        from docx import Document

        d = Document()
        d.add_paragraph(f"Confidential {MARKER} salary 123456")
        d.add_paragraph("<script>alert(1)</script> second paragraph")
        path = tmp / "secret-name-report2.docx"
        d.save(path)
        raw, events = stream(client, path)
        chunks = raw.split("event: result")[0].split("\n\n")
        before_result = "\n".join(c for c in chunks if c.startswith("event: stage"))  # the start event only lists stage names
        assert MARKER not in before_result and "salary" not in before_result
        assert "secret-name-report2" not in before_result and "uploads" not in before_result.lower()
        assert ".docx" not in before_result and "Users" not in before_result and "script" not in before_result.lower()
        for s in stage_events(events):
            assert set(s) <= {"id", "state", "seq", "elapsed_ms", "detail"}
            assert set(s.get("detail", {})) <= WHITELIST

    def test_detail_cleaning_drops_anything_not_on_the_whitelist(self):
        dirty = {"findings": 3, "text": "secret", "path": "C:/x", "reason": "my password is x", "status": "verified",
                 "by_severity": {"high": 2, "<script>": 1, "medium": "9"}, "codes": {"missing_text": 1, "a b c": 2},
                 "error_type": "bad type!", "issues": True, "blocks": -1}
        assert _clean_detail(dirty) == {"findings": 3, "status": "verified", "by_severity": {"high": 2},
                                        "codes": {"missing_text": 1}}

    def test_unexpected_failures_do_not_expose_server_internals(self, client, docx, monkeypatch):
        import services.parse_service as ps

        monkeypatch.setattr(ps, "route_regions", lambda *a, **k: (_ for _ in ()).throw(RuntimeError("C:/secret/path.key")))
        raw, events = stream(client, docx)
        assert events[-1][0] == "error" and "secret" not in raw and events[-1][1]["code"] == "INTERNAL_ERROR"


# ================================================================== tracker logic


class _Issue:
    def __init__(self, sev, code):
        self.severity, self.code = sev, code


class TestTracker:
    def _tracker(self):
        out = []
        return Tracker(out.append), out

    def test_a_multi_layer_stage_runs_once_and_closes_on_its_last_layer(self):
        t, out = self._tracker()
        for layer in ("integrity", "completeness"):
            t.engine(layer, "running")
            t.engine(layer, "completed")
        assert [(e["id"], e["state"]) for e in out] == [("checks", "running"), ("checks", "completed")]

    def test_issues_make_a_warning_and_are_counted_by_severity_and_code(self):
        t, out = self._tracker()
        t.engine("content", "running")
        t.engine("content", "completed", issues=[_Issue("high", "missing_text"), _Issue("low", "x")])
        assert out[-1]["state"] == "warning" and out[-1]["detail"]["by_severity"] == {"high": 1, "low": 1}

    def test_unreached_layers_are_skipped_not_faked(self):
        t, out = self._tracker()
        t.engine("inventory", "running")
        t.engine("inventory", "failed", error_type="X")
        t.finish_engine()
        final = {e["id"]: e for e in out}
        assert final["inventory"]["state"] == "failed" and final["content"]["state"] == "skipped"

    def test_a_cancelled_job_stops_at_its_next_report(self):
        t, _ = self._tracker()
        progress.install(t)
        t.cancelled.set()
        with pytest.raises(progress.Cancelled):
            progress.report("security", "running")
        progress.install(None)

    def test_reporting_is_a_noop_without_a_listener(self):
        progress.install(None)
        progress.report("security", "running")
        progress.report_engine("content", "running")


# ================================================================== cancellation and disconnects


class TestCancellation:
    def test_closing_the_stream_cancels_the_job_and_flags_the_worker(self):
        state = {}

        async def job():
            state["tracker"] = progress.current()
            try:
                await asyncio.sleep(30)
            except asyncio.CancelledError:
                state["cancelled"] = True
                raise

        async def run():
            gen = event_stream(job)
            assert (await gen.__anext__()).startswith(b"event: start")
            await asyncio.sleep(0.05)
            await gen.aclose()
            await asyncio.sleep(0.05)

        asyncio.run(run())
        assert state.get("cancelled") and state["tracker"].cancelled.is_set()

    def test_an_idle_stream_sends_keepalive_pings(self):
        async def job():
            await asyncio.sleep(0.5)
            return {"ok": 1}

        async def run():
            return [chunk async for chunk in event_stream(job, ping_seconds=0.1)]

        chunks = asyncio.run(run())
        assert any(c.startswith(b": ping") for c in chunks) and chunks[-1].startswith(b"event: result")


# ================================================================== authenticated streaming upload


class TestAuthenticatedStream:
    def test_requires_authentication(self, docx):
        import main

        main.app.dependency_overrides.clear()
        with TestClient(main.app) as c:
            with open(docx, "rb") as fh:
                r = c.post("/api/documents/stream", files={"file": ("a.docx", fh, "application/octet-stream")})
            assert r.status_code == 401

    def test_saves_to_history_like_the_plain_upload_and_returns_its_id(self, env, docx):
        client, fake, storage, who = env
        _, events = stream(client, docx, "/api/documents/stream")
        body = events[-1][1]
        assert events[-1][0] == "result" and body["document_id"] and body["result"]["validation"]
        saved = client.get(f"/api/documents/{body['document_id']}/result").json()["result"]
        assert saved["validation"] == body["result"]["validation"]
        assert storage, "the original was stored encrypted"

    def test_each_file_gets_its_own_document_id(self, env, docx, tmp):
        client, *_ = env
        _, a = stream(client, docx, "/api/documents/stream")
        _, b = stream(client, make_pdf(tmp / "b.pdf", pages=1), "/api/documents/stream")
        assert a[-1][1]["document_id"] != b[-1][1]["document_id"]
        assert a[-1][1]["result"]["file_type"] == "docx" and b[-1][1]["result"]["file_type"] == "pdf"

    def test_a_cancelled_upload_is_marked_failed_not_left_processing(self, env, docx, monkeypatch):
        client, fake, storage, who = env
        import routers.documents as docs

        calls = []
        real = docs._mark_failed

        async def spy(doc_id, run_id, user_id, code, message):
            calls.append(code)
            await real(doc_id, run_id, user_id, code, message)

        monkeypatch.setattr(docs, "_mark_failed", spy)

        import time

        monkeypatch.setattr(docs, "parse_upload", lambda *a, **k: time.sleep(1.5))  # runs in the worker thread

        async def run():
            from io import BytesIO

            from fastapi import UploadFile

            up = UploadFile(file=BytesIO(docx.read_bytes()), filename="a.docx")
            gen = event_stream(lambda: docs._store_and_parse(up, who["user"]))
            await gen.__anext__()
            await asyncio.sleep(0.5)
            await gen.aclose()
            await asyncio.sleep(0.3)

        asyncio.run(run())
        assert calls == ["CANCELLED"]
