"""Diligence intelligence: contradictions, totals, missing references, packs, maturity, exports."""

from __future__ import annotations

import io
import sys
from datetime import date
from pathlib import Path

import pytest
import yaml

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from rag import diligence, exports, maturity, packs  # noqa: E402
from tests.test_rag_core import room, wait_ready  # noqa: E402,F401

ROOT = Path(__file__).resolve().parents[2]
DEMO = ROOT / "demo" / "project_falcon"
FACTS = yaml.safe_load((ROOT / "eval" / "golden_qa.yaml").read_text())["facts"] if (ROOT / "eval" / "golden_qa.yaml").exists() else {}


def fact(value, raw, scale=1e7, cur="INR", label="Revenue"):
    return {"value": value * scale, "raw": raw, "scale": scale, "currency": cur, "label": label}


class TestToleranceRules:
    def test_equal_within_half_percent_or_rounding(self):
        assert diligence._equal(fact(452.0, "452.0"), fact(452.0, "452"))
        assert diligence._equal(fact(452.0, "452.0"), fact(453.5, "453.5"))        # 0.33 % < 0.5 %
        assert diligence._equal(fact(12.0, "12"), fact(12.4, "12.4"))              # within rounding of the coarser figure (±0.5)
        assert not diligence._equal(fact(452.0, "452.0"), fact(480.0, "480.0"))     # 6.2 %
        assert not diligence._equal(fact(100.0, "100.0"), fact(101.0, "101.0"))     # 1 % > 0.5 %

    def test_format_value(self):
        assert diligence.fmt_value(480.0 * 1e7, 1e7, "INR", "crore", "480.0") == "₹480.0 Cr"
        assert diligence.fmt_value(2.5e6, 1e6, "USD", "million", "2.5") == "$2.5 M"


class TestExportsAreSafe:
    @pytest.mark.parametrize("cell", ["=HYPERLINK(\"http://evil\",\"x\")", "+1+1", "-2+3", "@SUM(A1)", "\t=1", "\r=1"])
    def test_formula_like_cells_are_neutralised(self, cell):
        assert exports.guard(cell).startswith("'")

    def test_normal_text_and_numbers_untouched(self):
        assert exports.guard("Revenue") == "Revenue" and exports.guard(5) == 5 and exports.guard("") == ""

    def test_csv_and_xlsx_contain_no_live_formulas(self):
        cols, rows = ["=cmd|' /C calc'!A0", "Q"], [["+SUM(1,1)", "ok"], ["@x", "-5"]]
        text = exports.csv_bytes(cols, rows).decode("utf-8-sig")
        assert "'=cmd" in text and "'+SUM" in text and "'@x" in text and "'-5" in text
        import openpyxl

        ws = openpyxl.load_workbook(io.BytesIO(exports.xlsx_bytes(cols, rows))).active
        values = [c.value for row in ws.iter_rows() for c in row]
        assert all(not (isinstance(v, str) and v.startswith(("=", "+", "-", "@"))) for v in values)
        assert all(c.data_type != "f" for row in ws.iter_rows() for c in row)


class TestDates:
    def test_formats(self):
        got = {d for d, _, _ in maturity._parse_dates("Due 30 September 2026; renewal on Mar 31, 2027 and 15/12/2026; 31 Feb 2026 is not a date.")}
        assert got == {date(2026, 9, 30), date(2027, 3, 31), date(2026, 12, 15)}


@pytest.mark.skipif(not DEMO.exists(), reason="demo data room not generated")
class TestFalconRoom:
    @pytest.fixture()
    def falcon(self, room):
        ws = room.post("/api/workspaces", json={"name": "F"}).json()["workspace_id"]
        files = [("files", (f.name, f.read_bytes())) for f in sorted(DEMO.iterdir())]
        room.post(f"/api/workspaces/{ws}/documents", files=files)
        wait_ready(room, ws, len(files))
        return room, ws

    def test_contradictions_are_found_with_both_sides(self, falcon):
        room, ws = falcon
        cs = {(c["concept_label"], c["period"]): c for c in room.get(f"/api/workspaces/{ws}/contradictions").json()}
        rev = cs[("Revenue", "FY2024")]
        assert rev["gap_pct"] == pytest.approx((480.0 - 452.0) / 452.0 * 100, abs=0.06) and rev["severity"] == "high"
        assert rev["primary"]["high"]["filename"] == "Falcon_CIM.pdf" and rev["primary"]["low"]["filename"].startswith("Falcon_Audited")
        assert rev["primary"]["low"]["printed_page"] == "F-2" and rev["primary"]["high"]["bbox"]
        assert ("EBITDA", "FY2024") in cs and "adjusted" in cs[("EBITDA", "FY2024")]["note"].lower() or cs[("EBITDA", "FY2024")]["note"] == ""
        assert ("Total debt", "FY2025") not in cs          # the figures agree everywhere: no false alarm
        assert ("Revenue", "FY2025") not in cs

    def test_total_mismatch_is_detected(self, falcon):
        room, ws = falcon
        issues = diligence.total_checks(ws)
        assert len(issues) == 1 and issues[0]["filename"] == "Falcon_Management_Accounts.xlsx"
        assert issues[0]["diff"] == pytest.approx(-FACTS["monthly_total_typed_error"])
        assert not any(i["filename"] == "Falcon_Debt_Schedule.pdf" for i in issues)   # the 50-row two-page schedule adds up

    def test_missing_schedule_reference(self, falcon):
        room, ws = falcon
        assert [(m["label"], m["filename"]) for m in diligence.missing_references(ws)] == [("Schedule 3", "Falcon_Loan_Agreement.docx")]

    def test_maturity_wall_matches_ground_truth_and_has_receipts(self, falcon):
        room, ws = falcon
        data = room.get(f"/api/workspaces/{ws}/maturity-wall").json()
        wall = data["walls"][0]
        bars = {b["year"]: b for b in wall["bars"]}
        assert bars[2026]["value"] == pytest.approx(FACTS["maturing_2026"], abs=0.051) and bars[2027]["value"] == pytest.approx(FACTS["maturing_2027"], abs=0.051)
        assert sum(b["value"] for b in wall["bars"]) == pytest.approx(FACTS["total_debt"], abs=0.1)
        assert bars[2026]["receipt"]["operands"] and all(o["bbox"] for o in bars[2026]["receipt"]["operands"])
        events = data["events"]
        assert any(e["date"] == "2026-09-30" and e["kind"] == "maturity" for e in events)
        assert any("facilities mature" in e["label"] and e.get("count", 0) > 1 for e in events)   # same-day maturities are grouped

    def test_packs_matrix_and_export(self, falcon):
        room, ws = falcon
        fin = room.post(f"/api/workspaces/{ws}/packs/financials").json()
        rev = next(r for r in fin["rows"] if r["label"] == "Revenue")
        cells = {c["doc_id"]: c for c in rev["cells"]}
        names = {c["doc_id"]: c["filename"] for c in fin["columns"]}
        by_name = {names[d]: c for d, c in cells.items()}
        assert by_name["Falcon_CIM.pdf"]["text"].startswith("₹480.0 Cr") and by_name["Falcon_Audited_Financials_FY2025.pdf"]["text"].startswith("₹546.0 Cr")
        assert all(c["citations"] and c["citations"][0]["bbox"] for c in cells.values() if c["status"] == "found")
        contracts = room.post(f"/api/workspaces/{ws}/packs/contracts").json()
        assert {r["label"]: r["cells"][0]["status"] for r in contracts["rows"]}["Governing law"] == "found"
        assert {r["label"]: r["cells"][0]["status"] for r in contracts["rows"]}["Assignment"] == "not_found"
        r = room.get(f"/api/workspaces/{ws}/packs/financials/export?format=xlsx")
        assert r.status_code == 200 and r.content[:2] == b"PK"

    def test_seller_questions_are_cited_ranked_and_exportable(self, falcon):
        room, ws = falcon
        for p in ("financials", "contracts", "debt"):
            room.post(f"/api/workspaces/{ws}/packs/{p}")
        items = room.get(f"/api/workspaces/{ws}/seller-questions").json()
        kinds = {i["kind"] for i in items}
        assert {"contradiction", "total_mismatch", "missing_reference", "hardcoded_value", "hidden_instruction", "hidden_content", "pack_gap"} <= kinds
        assert [i["severity"] for i in items] == sorted((i["severity"] for i in items), key=["high", "medium", "low"].index)
        assert [i["n"] for i in items] == list(range(1, len(items) + 1))
        assert all(i["evidence"] for i in items if i["kind"] != "pack_gap")
        assert sum(1 for i in items if i["kind"] == "hidden_instruction") == 1                # one item, not one per finding
        edited = [{**items[0], "question": "=HYPERLINK(\"http://evil\",\"x\")"}] + items[1:3]
        for fmt, magic in (("docx", b"PK"), ("md", b"#"), ("csv", b"")):
            r = room.post(f"/api/workspaces/{ws}/seller-questions/export", json={"format": fmt, "items": edited})
            assert r.status_code == 200 and r.content.lstrip(b"\xef\xbb\xbf").startswith(magic or b"#")
        csv = room.post(f"/api/workspaces/{ws}/seller-questions/export", json={"format": "csv", "items": edited}).content.decode("utf-8-sig")
        assert "'=HYPERLINK" in csv and ",=HYPERLINK" not in csv
