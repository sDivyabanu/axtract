"""Generators for small test documents (no binary fixtures are checked in).

Each builder writes a file and returns its path. Ground truth is returned/declared next to the
builder so tests assert against the values the document was built from.
"""

from __future__ import annotations

import io
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402

BAR_LABELS = ["Q1", "Q2", "Q3", "Q4"]
BAR_VALUES = [120, 145, 160, 190]
LINE_X = [1, 2, 3, 4, 5]
LINE_Y = [2.0, 2.5, 3.1, 3.0, 4.2]
SCATTER_X = [10, 20, 30, 40]
SCATTER_Y = [5, 9, 14, 22]


# ---------------------------------------------------------------------------
# pictures
# ---------------------------------------------------------------------------


def bar_chart_png(path: Path, labels=BAR_LABELS, values=BAR_VALUES, data_labels=True) -> Path:
    fig, ax = plt.subplots(figsize=(8, 5), dpi=200)
    bars = ax.bar(labels, values, color="#2b6cb0")
    ax.set_title("Quarterly Revenue ($M)")
    ax.set_xlabel("Quarter")
    ax.set_ylabel("Revenue ($M)")
    if data_labels:
        for b, v in zip(bars, values):
            ax.text(b.get_x() + b.get_width() / 2, v + 2, str(v), ha="center")
    fig.savefig(path, facecolor="white")
    plt.close(fig)
    return path


def line_chart_png(path: Path) -> Path:
    fig, ax = plt.subplots(figsize=(8, 4.5), dpi=200)
    ax.plot(["Jan", "Feb", "Mar", "Apr", "May"], LINE_Y, marker="o", color="#c05621")
    ax.set_title("Users (millions)")
    fig.savefig(path, facecolor="white")
    plt.close(fig)
    return path


def equation_png(path: Path, latex: str = r"$E = mc^{2}$", size=(5, 1.6)) -> Path:
    fig = plt.figure(figsize=size, dpi=200)
    fig.text(0.5, 0.5, latex, fontsize=34, ha="center", va="center")
    fig.savefig(path, facecolor="white")
    plt.close(fig)
    return path


def text_png(path: Path, text: str = "Image OCR test text and some more words") -> Path:
    from PIL import Image, ImageDraw

    img = Image.new("RGB", (700, 160), "white")
    ImageDraw.Draw(img).text((20, 60), text, fill="black")
    img.save(path)
    return path


# ---------------------------------------------------------------------------
# PDFs
# ---------------------------------------------------------------------------


def vector_chart_pdf(path: Path) -> Path:
    """matplotlib PDF: bars, line and scatter drawn as paths (no images)."""
    fig, axs = plt.subplots(1, 3, figsize=(12, 3.5))
    axs[0].bar(BAR_LABELS, BAR_VALUES, color="#2b6cb0")
    axs[0].set_title("Revenue")
    axs[1].plot(LINE_X, LINE_Y, marker="o", color="#c05621")
    axs[1].set_title("Users")
    axs[2].scatter(SCATTER_X, SCATTER_Y, color="#2f855a")
    axs[2].set_title("Scatter")
    fig.savefig(path)
    plt.close(fig)
    return path


def pdf_with_images(path: Path, images: list[tuple[Path, str]], intro: str = "") -> Path:
    """reportlab PDF with a caption line above each embedded image."""
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.utils import ImageReader
    from reportlab.pdfgen import canvas

    c = canvas.Canvas(str(path), pagesize=A4)
    w, h = A4
    y = h - 60
    if intro:
        c.setFont("Helvetica", 11)
        c.drawString(60, y, intro)
        y -= 30
    for img_path, caption in images:
        img = ImageReader(str(img_path))
        iw, ih = img.getSize()
        draw_w = min(420, w - 120)
        draw_h = draw_w * ih / iw
        if y - draw_h - 40 < 40:
            c.showPage()
            y = h - 60
        c.setFont("Helvetica-Bold", 14)
        c.drawString(60, y, caption)
        y -= 14
        c.drawImage(img, 60, y - draw_h, draw_w, draw_h)
        y -= draw_h + 30
    c.save()
    return path


def pdf_with_paragraphs(path: Path, paragraphs: list[str]) -> Path:
    """reportlab PDF whose paragraphs wrap over several lines."""
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import Paragraph, SimpleDocTemplate

    styles = getSampleStyleSheet()
    SimpleDocTemplate(str(path), pagesize=A4).build([Paragraph(p, styles["BodyText"]) for p in paragraphs])
    return path


# ---------------------------------------------------------------------------
# Office documents
# ---------------------------------------------------------------------------

OMML_NS = 'xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math"'


def _omml_fraction_equation() -> str:
    """x = (-b ± sqrt(b^2 - 4ac)) / 2a as OMML."""

    def r(t):
        return f"<m:r><m:t>{t}</m:t></m:r>"

    return (
        f"<m:oMathPara {OMML_NS}><m:oMath>{r('x=')}<m:f><m:num>{r('−b±')}<m:rad><m:radPr><m:degHide m:val=\"1\"/></m:radPr>"
        f"<m:deg/><m:e><m:sSup><m:e>{r('b')}</m:e><m:sup>{r('2')}</m:sup></m:sSup>{r('−4ac')}</m:e></m:rad></m:num>"
        f"<m:den>{r('2a')}</m:den></m:f></m:oMath></m:oMathPara>"
    )


OMML_EXPECTED = r"x=\frac{-b\pm \sqrt{{b}^{2}-4ac}}{2a}"


def docx_with_omml(path: Path) -> Path:
    import docx
    from lxml import etree

    d = docx.Document()
    d.add_heading("Math report", 1)
    d.add_paragraph("The quadratic formula follows.")
    p = d.add_paragraph()
    p._p.append(etree.fromstring(_omml_fraction_equation()))
    d.add_paragraph("End of math.")
    d.save(path)
    return path


def docx_with_picture(path: Path, picture: Path, caption: str = "Figure") -> Path:
    import docx
    from docx.shared import Inches

    d = docx.Document()
    d.add_heading("Picture report", 1)
    d.add_paragraph(caption)
    d.add_picture(str(picture), width=Inches(5))
    d.add_paragraph("After the picture.")
    d.save(path)
    return path


def _chart_xml_from_pptx() -> bytes:
    """A real chart part (clustered column) produced by python-pptx's chart XML writer."""
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.chart.xmlwriter import ChartXmlWriter

    cd = CategoryChartData()
    cd.categories = BAR_LABELS
    cd.add_series("Revenue", BAR_VALUES)
    return ChartXmlWriter(XL_CHART_TYPE.COLUMN_CLUSTERED, cd).xml.encode("utf-8")


def docx_with_native_chart(path: Path) -> Path:
    """DOCX carrying a native chart part (word/charts/chart1.xml) referenced from a drawing."""
    import docx
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    from docx.opc.packuri import PackURI
    from docx.opc.part import Part
    from lxml import etree

    d = docx.Document()
    d.add_heading("Chart report", 1)
    part = Part(
        PackURI("/word/charts/chart1.xml"),
        "application/vnd.openxmlformats-officedocument.drawingml.chart+xml",
        _chart_xml_from_pptx(),
        d.part.package,
    )
    rid = d.part.relate_to(part, RT.CHART)
    p = d.add_paragraph()
    drawing = etree.fromstring(
        f"""<w:drawing xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"
            xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"
            xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"
            xmlns:c="http://schemas.openxmlformats.org/drawingml/2006/chart"
            xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">
          <wp:inline><wp:extent cx="4572000" cy="2743200"/><wp:docPr id="1" name="Chart 1"/>
            <a:graphic><a:graphicData uri="http://schemas.openxmlformats.org/drawingml/2006/chart">
              <c:chart r:id="{rid}"/></a:graphicData></a:graphic></wp:inline></w:drawing>"""
    )
    run = p.add_run()
    run._r.append(drawing)
    d.add_paragraph("After the chart.")
    d.save(path)
    return path


def pptx_with_native_chart(path: Path) -> Path:
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.util import Inches

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = "Native chart"
    cd = CategoryChartData()
    cd.categories = BAR_LABELS
    cd.add_series("Revenue", BAR_VALUES)
    slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(1.8), Inches(7), Inches(4.5), cd)
    prs.save(path)
    return path


def pptx_with_picture(path: Path, picture: Path, title: str = "Picture") -> Path:
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = title
    slide.shapes.add_picture(str(picture), Inches(1), Inches(1.8), width=Inches(7))
    prs.save(path)
    return path


def pptx_with_omml(path: Path) -> Path:
    """Text box whose paragraph holds an OMML equation inside mc:AlternateContent / a14:m."""
    from lxml import etree
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = "Math slide"
    tb = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(7), Inches(1.5))
    p = tb.text_frame.paragraphs[0]._p
    alt = etree.fromstring(
        f"""<mc:AlternateContent xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006"
              xmlns:a14="http://schemas.microsoft.com/office/drawing/2010/main">
            <mc:Choice Requires="a14"><a14:m>{_omml_fraction_equation()}</a14:m></mc:Choice>
            <mc:Fallback><a:r xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"><a:t>fallback picture</a:t></a:r></mc:Fallback>
          </mc:AlternateContent>"""
    )
    p.append(alt)
    prs.save(path)
    return path


def xlsx_with_merges(path: Path) -> Path:
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet"
    ws["A1"] = "Title"
    ws.merge_cells("A1:C1")
    ws.append(["Name", "H1", "H2"])
    ws.append(["Alpha", 1, 2])
    wb.save(path)
    return path


def png_bytes(img_path: Path) -> bytes:
    buf = io.BytesIO()
    from PIL import Image

    Image.open(img_path).save(buf, format="PNG")
    return buf.getvalue()
