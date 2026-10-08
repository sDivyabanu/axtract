"""AXTRACT Verify layer 2: completeness, source inventory vs AXTRACT's real extraction.

Every case uses a genuine file parsed by the real pipeline. "Silent loss" is simulated by removing
blocks from that real extraction and checking that Verify notices.
"""

from __future__ import annotations

from collections import Counter
from dataclasses import dataclass
from pathlib import Path

import pytest

from models.document import BlockType as T, DocumentBlock
from tests.verify_fixtures import drop, extract, make_docx, make_pdf, make_pptx, make_xlsx, png_file
from verify.completeness import MatchStatus, check_completeness
from verify.ids import IdSequence
from verify.inventory import build_inventory
from verify.integrity import run_integrity
from verify.models import (
    CheckName, CheckOutcome, CheckResult, EvidenceKind, Severity, ValidationReport, ValidationStatus,
)
from verify.rollup import UnitChecks, build_report, merge_units

SAMPLES = Path(__file__).resolve().parents[2] / "sample_files"


@dataclass
class Case:
    path: Path
    ext: str
    resp: object
    inv: object

    def run(self, resp=None):
        return check_completeness(self.inv, resp or self.resp)


def make_case(path: Path) -> Case:
    ext = path.suffix.lstrip(".")
    return Case(path, ext, extract(path), build_inventory(path, ext))


@pytest.fixture(scope="module")
def tmp(tmp_path_factory):
    return tmp_path_factory.mktemp("completeness")


@pytest.fixture(scope="module")
def xlsx(tmp):
    return make_case(make_xlsx(tmp / "book.xlsx"))


@pytest.fixture(scope="module")
def xlsx_plain(tmp):
    return make_case(make_xlsx(tmp / "plain.xlsx", chart=False, picture=False))


@pytest.fixture(scope="module")
def pptx(tmp):
    return make_case(make_pptx(tmp / "deck.pptx"))


@pytest.fixture(scope="module")
def docx(tmp):
    return make_case(make_docx(tmp / "doc.docx"))


@pytest.fixture(scope="module")
def pdf(tmp):
    return make_case(make_pdf(tmp / "doc.pdf", pages=3, scanned_page=3))


def codes(out) -> Counter:
    return Counter(i.code for i in out.issues)


def issue(out, code, unit=None):
    found = [i for i in out.issues if i.code == code and (unit is None or i.locator.unit.key == unit)]
    assert found, f"{code} on {unit} not raised; got {[(i.code, i.locator.unit.key) for i in out.issues]}"
    return found[0]


def by_name(name):
    return lambda b: b.metadata.get("shape_name") == name


def of_type(t, page=None):
    return lambda b: b.type == t and (page is None or b.page == page)


def with_text(fragment):
    return lambda b: fragment in b.content


def block_of(resp, pred):
    return next(b for b in resp.blocks if pred(b))


def clone(case, mutate) -> object:
    r = case.resp.model_copy(deep=True)
    mutate(r)
    return r


# ====================================================================================== XLSX


class TestXlsxCompleteness:
    def test_a_workbook_without_drawings_is_complete(self, xlsx_plain):
        out = xlsx_plain.run()
        assert out.issues == [] and out.stats["unmatched"] == 0

    def test_charts_and_pictures_are_flagged_because_the_extractor_does_not_read_them(self, xlsx):
        out = xlsx.run()
        assert codes(out) == {"possible_missing_chart": 1, "possible_missing_image": 1}
        c = issue(out, "possible_missing_chart", "sheet:1")
        assert "Sales chart" in c.message and c.evidence[0].detail["chart_types"] == ["barChart"]
        assert c.evidence[0].kind == EvidenceKind.DETERMINISTIC and c.evidence[0].engine.startswith("ooxml:")

    def test_formulas_are_evidence_but_never_reported_missing(self, xlsx_plain):
        m = next(m for m in xlsx_plain.run().matches if m.object_id.endswith("formulas"))
        assert m.status == MatchStatus.NOT_EXPECTED and "cached values" in " ".join(m.reasons)

    def test_hidden_and_empty_sheets(self, xlsx_plain):
        out = xlsx_plain.run()
        hidden = next(m for m in out.matches if m.object_id == "xlsx:sheet3:cells")
        assert hidden.status == MatchStatus.MATCHED  # hidden sheet content is extracted
        assert not any(m.unit == "sheet:4" and m.status == MatchStatus.UNMATCHED for m in out.matches)  # empty sheet: nothing to lose

    def test_missing_sheet(self, xlsx_plain):
        out = xlsx_plain.run(drop(xlsx_plain.resp, lambda b: b.metadata.get("sheet_name") == "Costs"))
        i = issue(out, "possible_missing_sheet", "sheet:2")
        assert i.severity == Severity.HIGH and "Costs" in i.message

    def test_a_dropped_hidden_sheet_is_still_missing_content(self, xlsx_plain):
        out = xlsx_plain.run(drop(xlsx_plain.resp, lambda b: b.metadata.get("sheet_name") == "Hidden notes"))
        i = issue(out, "possible_missing_sheet", "sheet:3")
        assert i.evidence[0].detail["sheet_state"] == "hidden"

    def test_missing_cells_one_cell_is_medium_and_names_the_cell(self, xlsx_plain):
        def mutate(r):
            blk = block_of(r, lambda b: b.metadata.get("sheet_name") == "Sales")
            blk.metadata["rows"][1][1] = ""  # B2
        out = xlsx_plain.run(clone(xlsx_plain, mutate))
        i = issue(out, "possible_missing_cells", "sheet:1")
        assert i.severity == Severity.MEDIUM and i.evidence[0].detail["first_missing"] == ["B2"]
        assert i.evidence[0].detail["missing"] == 1 and i.evidence[0].detail["populated"] == 14

    def test_losing_many_cells_is_high(self, xlsx_plain):
        def mutate(r):
            blk = block_of(r, lambda b: b.metadata.get("sheet_name") == "Sales")
            for row in blk.metadata["rows"][1:]:
                row[1:] = [""] * (len(row) - 1)
        i = issue(xlsx_plain.run(clone(xlsx_plain, mutate)), "possible_missing_cells", "sheet:1")
        assert i.severity == Severity.HIGH and i.evidence[0].detail["missing_ratio"] >= 0.1

    def test_a_missing_cell_inside_an_excel_table_also_reports_the_table(self, xlsx_plain):
        def mutate(r):
            block_of(r, lambda b: b.metadata.get("sheet_name") == "Sales").metadata["rows"][2][2] = None  # C3, inside A1:C4
        out = xlsx_plain.run(clone(xlsx_plain, mutate))
        assert {"possible_missing_cells", "possible_missing_table"} <= set(codes(out))
        assert issue(out, "possible_missing_table").evidence[0].detail["missing"] == 1

    def test_a_missing_cell_outside_every_excel_table_does_not_report_the_table(self, xlsx_plain):
        def mutate(r):
            block_of(r, lambda b: b.metadata.get("sheet_name") == "Sales").metadata["rows"][0][5] = ""  # F1, outside A1:C4
        out = xlsx_plain.run(clone(xlsx_plain, mutate))
        assert "possible_missing_cells" in codes(out) and "possible_missing_table" not in codes(out)

    def test_merged_ranges_present_are_matched_and_missing_ones_flagged(self, xlsx_plain):
        base = next(m for m in xlsx_plain.run().matches if m.object_type.value == "merged_range")
        assert base.status == MatchStatus.MATCHED and "F1:G2" in base.reasons[0]

        def mutate(r):
            block_of(r, lambda b: b.metadata.get("sheet_name") == "Sales").metadata["merged_cells"] = {"has_merged_cells": False, "merged_regions": []}
        i = issue(xlsx_plain.run(clone(xlsx_plain, mutate)), "possible_missing_merged_range", "sheet:1")
        assert i.evidence[0].source == "F1:G2" and i.evidence[0].detail["expected"] == [0, 5, 2, 2]

    def test_merged_ranges_are_not_verifiable_when_the_extractor_says_it_did_not_read_them(self, xlsx_plain):
        def mutate(r):
            blk = block_of(r, lambda b: b.metadata.get("sheet_name") == "Sales")
            blk.metadata["merged_cells"] = {"has_merged_cells": False, "merged_regions": []}
            blk.metadata["flags"] = ["merged_cells_not_read_large_file"]
        out = xlsx_plain.run(clone(xlsx_plain, mutate))
        assert "possible_missing_merged_range" not in codes(out)
        assert next(m for m in out.matches if m.object_type.value == "merged_range").status == MatchStatus.NOT_VERIFIABLE

    def test_a_chart_block_resolves_the_chart_but_one_block_cannot_resolve_both(self, xlsx):
        def add_one(r):
            r.blocks.append(DocumentBlock(id="x-fig", type=T.FIGURE, content="[image]", page=1, extractor="test", reading_order=99))
        out = xlsx.run(clone(xlsx, add_one))
        assert sum(codes(out).values()) == 1  # one source object (chart or picture) is still unrepresented

        def add_two(r):
            add_one(r)
            r.blocks.append(DocumentBlock(id="x-chart", type=T.CHART, content="c", page=1, extractor="test", reading_order=100))
        assert xlsx.run(clone(xlsx, add_two)).issues == []

    def test_every_matched_object_records_why(self, xlsx_plain):
        for m in xlsx_plain.run().matches:
            if m.status == MatchStatus.MATCHED:
                assert m.reasons and m.block_ids and m.engine.startswith("ooxml:") and m.evidence_kind == EvidenceKind.DETERMINISTIC

    def test_unit_checks_distinguish_deterministic_evidence(self, xlsx):
        out = xlsx.run()
        sales = next(u for u in out.units if u.unit.key == "sheet:1")
        costs = next(u for u in out.units if u.unit.key == "sheet:2")
        assert sales.checks[0].outcome == CheckOutcome.FAILED and costs.checks[0].outcome == CheckOutcome.PASSED
        assert costs.checks[0].evidence_kind == EvidenceKind.DETERMINISTIC and "presence only" in costs.checks[0].summary


# ====================================================================================== PPTX


class TestPptxCompleteness:
    def test_a_complete_extraction_raises_nothing_and_records_why_each_object_matched(self, pptx):
        out = pptx.run()
        assert out.issues == []
        why = {m.object_id: m for m in out.matches if m.status == MatchStatus.MATCHED}
        box = next(m for m in why.values() if "shape name equal" in m.reasons)
        assert box.block_ids and any(r.startswith("text_recall=") for r in box.reasons)
        tbl = next(m for m in why.values() if m.object_type.value == "table")
        assert any("bbox_iou" in r for r in tbl.reasons) and any("dimensions 3x3 equal" in r for r in tbl.reasons)

    def test_missing_text_shape(self, pptx):
        i = issue(pptx.run(drop(pptx.resp, by_name("Summary box"))), "possible_missing_text", "slide:1")
        assert "Revenue grew" in i.evidence[0].source and i.evidence[0].engine.startswith("ooxml:ppt/slides")

    def test_partially_lost_text_reports_the_missing_words(self, pptx):
        def mutate(r):
            block_of(r, by_name("Summary box")).content = "Revenue grew"
        i = issue(pptx.run(clone(pptx, mutate)), "possible_missing_text", "slide:1")
        assert "strongly" in i.evidence[0].detail["missing_words"] and i.evidence[0].detail["recall"] < 0.85

    def test_missing_title(self, pptx):
        out = pptx.run(drop(pptx.resp, lambda b: b.page == 1 and b.type == T.HEADING))
        assert "possible_missing_heading" in codes(out)

    def test_missing_table(self, pptx):
        i = issue(pptx.run(drop(pptx.resp, of_type(T.TABLE))), "possible_missing_table", "slide:1")
        assert "3x3" in i.message

    def test_missing_picture(self, pptx):
        i = issue(pptx.run(drop(pptx.resp, by_name("Team photo"))), "possible_missing_image", "slide:1")
        assert "Team photo" in i.message

    def test_missing_chart(self, pptx):
        i = issue(pptx.run(drop(pptx.resp, of_type(T.CHART))), "possible_missing_chart", "slide:2")
        assert i.evidence[0].detail["series_count"] == 1

    def test_text_inside_a_group_is_expected_and_its_loss_is_detected(self, pptx):
        out = pptx.run(drop(pptx.resp, by_name("Group note")))
        assert issue(out, "possible_missing_text", "slide:2").evidence[0].source == "Text that lives inside a group shape"
        group = next(m for m in pptx.run().matches if m.object_type.value == "group")
        assert group.status == MatchStatus.NOT_EXPECTED  # the container itself is not content

    def test_a_slide_that_lost_everything_is_one_clear_issue_not_a_flood(self, pptx):
        out = pptx.run(drop(pptx.resp, lambda b: b.page == 2))
        assert codes(out) == {"possible_missing_slide_content": 1}
        assert issue(out, "possible_missing_slide_content", "slide:2").severity == Severity.HIGH

    def test_fewer_slides_reported_than_exist(self, pptx):
        out = pptx.run(clone(pptx, lambda r: setattr(r, "page_count", 1)))
        assert issue(out, "unit_count_mismatch", "document:1").evidence[0].source == "2 slides"

    def test_speaker_notes_are_not_reported_missing(self, pptx):
        notes = [m for m in pptx.run().matches if m.status == MatchStatus.NOT_EXPECTED and "notes" in " ".join(m.reasons)]
        assert notes and not any("note" in i.message.lower() for i in pptx.run().issues)

    def test_matching_survives_a_missing_shape_name_by_falling_back_to_geometry(self, pptx):
        def mutate(r):
            block_of(r, by_name("Summary box")).metadata.pop("shape_name")
        out = pptx.run(clone(pptx, mutate))
        assert "possible_missing_text" not in codes(out)
        m = next(m for m in out.matches if m.object_id.endswith("obj3") or "inside shape bbox" in " ".join(m.reasons))
        assert any("inside shape bbox" in r for r in m.reasons)


# ====================================================================================== DOCX


class TestDocxCompleteness:
    def test_a_complete_extraction_raises_nothing(self, docx):
        out = docx.run()
        assert out.issues == [] and out.stats["unmatched"] == 0

    def test_headers_are_outside_the_contract(self, docx):
        h = next(m for m in docx.run().matches if m.status == MatchStatus.NOT_EXPECTED)
        assert "headers and footers" in h.reasons[0]

    def test_missing_paragraph(self, docx):
        i = issue(docx.run(drop(docx.resp, with_text("operating costs remained flat"))), "possible_missing_text")
        assert "operating costs" in i.evidence[0].source and i.evidence[0].detail["path"].startswith("body[")

    def test_missing_heading(self, docx):
        assert "possible_missing_heading" in codes(docx.run(drop(docx.resp, with_text("Financial Results"))))

    def test_missing_list_item(self, docx):
        i = issue(docx.run(drop(docx.resp, with_text("Second list entry"))), "possible_missing_list")
        assert "list item" in i.message

    def test_missing_table(self, docx):
        assert "possible_missing_table" in codes(docx.run(drop(docx.resp, of_type(T.TABLE))))

    def test_missing_picture(self, docx):
        assert "possible_missing_image" in codes(docx.run(drop(docx.resp, of_type(T.FIGURE))))

    def test_missing_equation(self, docx):
        assert "possible_missing_equation" in codes(docx.run(drop(docx.resp, of_type(T.EQUATION))))

    def test_no_page_or_bbox_is_invented_for_issues(self, docx):
        i = issue(docx.run(drop(docx.resp, with_text("Financial Results"))), "possible_missing_heading")
        assert i.locator.bbox is None and i.locator.unit.key == "section:1"

    def test_small_wording_differences_do_not_count_as_missing(self, docx):
        def mutate(r):
            block_of(r, with_text("expanded into three")).content = "The company expanded in three new markets during the year"
        assert docx.run(clone(docx, mutate)).issues == []

    def test_a_heading_extracted_as_a_paragraph_is_still_present(self, docx):
        def mutate(r):
            block_of(r, with_text("Annual Report")).type = T.PARAGRAPH
        out = docx.run(clone(docx, mutate))
        assert out.issues == []
        m = next(m for m in out.matches if m.object_type.value == "heading" and "type differs" in " ".join(m.reasons))
        assert m.status == MatchStatus.MATCHED

    def test_matching_is_one_to_one_for_identical_paragraphs(self, tmp):
        import docx as d

        doc = d.Document()
        doc.add_paragraph("This sentence is repeated twice in the document body.")
        doc.add_paragraph("This sentence is repeated twice in the document body.")
        doc.save(tmp / "twice.docx")
        case = make_case(tmp / "twice.docx")
        assert case.run().issues == []
        out = case.run(drop(case.resp, lambda b: b.id == case.resp.blocks[0].id))  # lose ONE of the two
        assert codes(out) == {"possible_missing_text": 1}
        matched = [m for m in out.matches if m.status == MatchStatus.MATCHED]
        assert len(matched) == 1 and sum(len(m.block_ids) for m in matched) == 1

    def test_matching_is_one_to_one_for_identical_tables(self, tmp):
        import docx as d

        doc = d.Document()
        for _ in range(2):
            t = doc.add_table(rows=2, cols=2)
            for r in range(2):
                for c in range(2):
                    t.cell(r, c).text = f"same{r}{c}"
            doc.add_paragraph("separator paragraph between the tables")
        doc.save(tmp / "twotables.docx")
        case = make_case(tmp / "twotables.docx")
        assert case.run().issues == []
        out = case.run(drop(case.resp, lambda b: b.id == next(x for x in case.resp.blocks if x.type == T.TABLE).id))
        assert codes(out) == {"possible_missing_table": 1}

    def test_a_document_with_nothing_special_has_no_phantom_objects(self, tmp):
        import docx as d

        doc = d.Document()
        doc.add_paragraph("Only one plain paragraph lives here.")
        doc.save(tmp / "one.docx")
        case = make_case(tmp / "one.docx")
        assert case.run().issues == [] and case.run().stats["matched"] == 1


# ====================================================================================== PDF


class TestPdfCompleteness:
    def test_a_complete_multi_page_extraction_raises_nothing(self, pdf):
        out = pdf.run()
        assert out.issues == []
        assert {u.unit.key for u in out.units} == {"page:1", "page:2", "page:3"}

    def test_evidence_kinds_and_engines_are_stated_per_signal(self, pdf):
        out = pdf.run()
        kinds = {(m.object_type.value, m.evidence_kind, m.engine) for m in out.matches if m.status == MatchStatus.MATCHED}
        assert ("text", EvidenceKind.DETERMINISTIC, "pypdfium2") in kinds  # page text layer
        assert ("text", EvidenceKind.HEURISTIC, "pdfplumber") in kinds  # regions
        assert ("table", EvidenceKind.HEURISTIC, "pdfplumber") in kinds  # candidates
        assert ("picture", EvidenceKind.DETERMINISTIC, "pypdfium2") in kinds

    def test_a_table_match_says_the_candidate_is_not_ground_truth(self, pdf):
        m = next(m for m in pdf.run().matches if m.object_type.value == "table")
        assert m.status == MatchStatus.MATCHED and "candidate only: not ground truth" in m.reasons

    def test_missing_native_text_is_localised_to_a_region_with_a_bbox(self, pdf):
        out = pdf.run(drop(pdf.resp, with_text("Management expects continued growth")))
        i = issue(out, "possible_missing_text", "page:2" if False else None)
        assert i.locator.bbox and i.locator.page in (1, 2) and "Management expects" in i.evidence[0].source
        assert i.evidence[0].kind == EvidenceKind.HEURISTIC and i.evidence[0].detail["grouping_kind"] == "heuristic"
        assert i.evidence[0].detail["missing_words"]

    def test_losing_everything_on_one_page_flags_only_that_page(self, pdf):
        out = pdf.run(drop(pdf.resp, lambda b: b.page == 2))
        units = {i.locator.unit.key for i in out.issues}
        assert units == {"page:2"} and "possible_missing_text" in codes(out)

    def test_missing_table_candidate_is_medium_when_its_text_is_gone_too(self, pdf):
        i = issue(pdf.run(drop(pdf.resp, of_type(T.TABLE))), "possible_missing_table", "page:1")
        assert i.severity == Severity.MEDIUM and i.locator.bbox and i.evidence[0].kind == EvidenceKind.HEURISTIC
        assert i.evidence[0].detail["limitation"].startswith("ruling-line")

    def test_a_table_flattened_into_paragraphs_is_only_low_because_the_text_survives(self, pdf):
        def mutate(r):
            t = block_of(r, of_type(T.TABLE))
            t.type, t.metadata = T.PARAGRAPH, {}
        out = pdf.run(clone(pdf, mutate))
        i = issue(out, "possible_missing_table", "page:1")
        assert i.severity == Severity.LOW and i.evidence[0].detail["text_present_as_other_blocks"] is True

    def test_image_evidence_missing_figure(self, pdf):
        i = issue(pdf.run(drop(pdf.resp, of_type(T.FIGURE))), "possible_missing_image", "page:1")
        assert i.severity == Severity.MEDIUM and i.locator.bbox and "of the page" in i.evidence[0].source

    def test_a_scanned_page_is_not_verifiable_and_never_pretends_to_be(self, pdf):
        out = pdf.run()
        u3 = next(u for u in out.units if u.unit.key == "page:3")
        assert any(c.outcome == CheckOutcome.NOT_VERIFIABLE and "OCR" in c.summary for c in u3.checks)
        assert not any(c.outcome == CheckOutcome.PASSED for c in u3.checks)
        m = next(m for m in out.matches if m.status == MatchStatus.NOT_VERIFIABLE)
        assert "same engine" in m.reasons[0] or "engine AXTRACT uses" in m.reasons[0]

    def test_losing_a_scanned_page_cannot_be_detected_from_independent_evidence(self, pdf):
        out = pdf.run(drop(pdf.resp, lambda b: b.page == 3))
        assert not [i for i in out.issues if i.locator.unit.key == "page:3"]  # honest: we cannot know

    def test_the_full_page_scan_image_is_not_expected_as_a_figure(self, pdf):
        m = next(m for m in pdf.run().matches if m.unit == "page:3" and m.object_type.value == "picture")
        assert m.status == MatchStatus.NOT_EXPECTED and "scan layer" in m.reasons[0]

    def test_fewer_pages_reported_than_exist(self, pdf):
        out = pdf.run(clone(pdf, lambda r: setattr(r, "page_count", 2)))
        assert issue(out, "unit_count_mismatch", "document:1").evidence[0].source == "3 pages"

    def test_completeness_deliberately_does_not_judge_small_wording_changes(self, pdf):
        def mutate(r):
            b = block_of(r, with_text("Revenue increased strongly"))
            b.content = b.content.replace("(page 1)", "(page 7)")
        assert pdf.run(clone(pdf, mutate)).issues == []  # that is the content layer's job (next step)

    @pytest.mark.skipif(not (SAMPLES / "01_cross_page_table.pdf").exists(), reason="sample missing")
    def test_a_cross_page_merged_table_is_not_reported_missing_on_its_later_pages(self):
        case = make_case(SAMPLES / "01_cross_page_table.pdf")
        assert case.run().issues == []
        tables = [m for m in case.run().matches if m.object_type.value == "table"]
        assert tables and all(m.status == MatchStatus.MATCHED for m in tables)


# ====================================================================================== IMAGE


class TestImageCompleteness:
    def test_images_are_not_verifiable_not_passed(self, tmp):
        from PIL import ImageDraw, Image

        img = Image.new("RGB", (400, 120), "white")
        ImageDraw.Draw(img).text((20, 40), "Invoice total 1200", fill="black")
        img.save(tmp / "scan.png")
        case = make_case(tmp / "scan.png")
        out = case.run()
        assert out.issues == [] and out.stats == {"issues": 0, "source_objects": 1, "matched": 0, "unmatched": 0,
                                                    "not_expected": 0, "not_verifiable": 1}
        u = out.units[0]
        assert [c.outcome for c in u.checks] == [CheckOutcome.NOT_VERIFIABLE] and "engine AXTRACT used" in u.checks[0].summary
        assert case.inv.properties["width_px"] == 400

    def test_losing_all_ocr_text_cannot_be_detected_without_an_independent_reader(self, tmp):
        case = make_case(png_file(tmp / "blank.png", (200, 100)))
        assert case.run(drop(case.resp, lambda b: True)).issues == []


# ====================================================================================== cross-cutting


class TestCountsAloneNeverVerify:
    @pytest.mark.parametrize("fixture", ["pptx", "docx", "xlsx_plain", "pdf"])
    def test_a_perfect_completeness_result_is_still_not_verifiable(self, fixture, request):
        case = request.getfixturevalue(fixture)
        out = case.run()
        report = build_report(out.units, out.issues)
        assert report.status == ValidationStatus.NOT_VERIFIABLE
        assert ValidationStatus.VERIFIED not in {u.status for u in report.units}
        passing = [u for u in report.units if any(c.outcome == CheckOutcome.PASSED for c in u.checks)]
        assert passing and all("counts alone" in u.status_reasons[0] for u in passing)

    def test_verified_needs_a_content_bearing_check_on_top(self, pptx):
        out = pptx.run()
        extra = [UnitChecks(u.unit, [CheckResult(check=CheckName.CONTENT, outcome=CheckOutcome.PASSED,
                                                  evidence_kind=EvidenceKind.DETERMINISTIC, engine="content-layer")])
                 for u in out.units]
        report = build_report(merge_units(out.units, extra), out.issues)
        assert report.status == ValidationStatus.VERIFIED
        assert all(u.verified_by == ["content:content-layer(deterministic)"] for u in report.units)

    def test_a_completeness_problem_blocks_verified_even_with_content_evidence(self, pptx):
        out = pptx.run(drop(pptx.resp, by_name("Summary box")))
        extra = [UnitChecks(u.unit, [CheckResult(check=CheckName.CONTENT, outcome=CheckOutcome.PASSED,
                                                  evidence_kind=EvidenceKind.DETERMINISTIC, engine="content-layer")])
                 for u in out.units]
        report = build_report(merge_units(out.units, extra), out.issues)
        assert report.status == ValidationStatus.REVIEW_REQUIRED


class TestLayersCombine:
    def test_integrity_and_completeness_merge_into_one_consistent_report(self, pptx):
        ids = IdSequence("iss")
        integ = run_integrity(pptx.resp, ids)
        comp = pptx.run()
        comp2 = check_completeness(pptx.inv, drop(pptx.resp, by_name("Summary box")), ids)
        units = merge_units(integ.units, comp2.units)
        report = build_report(units, integ.issues + comp2.issues)
        assert report.status == ValidationStatus.REVIEW_REQUIRED
        s1 = next(u for u in report.units if u.unit.key == "slide:1")
        assert {c.check for c in s1.checks} == {CheckName.EXTRACTION_INTEGRITY, CheckName.COMPLETENESS}
        assert ValidationReport.model_validate_json(report.model_dump_json()) == report
        assert comp.issues == []

    def test_issue_ids_never_collide_when_the_sequence_is_shared(self, docx):
        ids = IdSequence("iss")
        a = check_completeness(docx.inv, drop(docx.resp, with_text("Financial Results")), ids)
        b = run_integrity(docx.resp.model_copy(update={"blocks": docx.resp.blocks[:0]}), ids)
        all_ids = [i.id for i in a.issues + b.issues]
        assert len(all_ids) == len(set(all_ids)) and a.issues


class TestIsolationAndHonesty:
    def test_an_unreadable_source_makes_completeness_not_verifiable_and_raises_nothing(self, tmp):
        (tmp / "broken.xlsx").write_bytes(b"nope")
        inv = build_inventory(tmp / "broken.xlsx", "xlsx")
        resp = make_case(make_xlsx(tmp / "ok.xlsx")).resp
        out = check_completeness(inv, resp)
        assert out.issues == [] and out.notes and "unavailable" in out.notes[0]
        assert [c.outcome for c in out.units[0].checks] == [CheckOutcome.NOT_VERIFIABLE]
        assert resp.blocks  # the extraction is untouched

    def test_a_crashing_comparator_is_contained(self, pptx, monkeypatch):
        import verify.completeness as cm

        def boom(ctx):
            raise RuntimeError("bug in comparator")
        monkeypatch.setitem(cm._COMPARATORS, "pptx", boom)
        out = pptx.run()
        assert out.issues == [] and "failed" in out.notes[0]
        assert out.units[0].checks[0].outcome == CheckOutcome.NOT_VERIFIABLE

    def test_unknown_formats_are_not_verifiable(self, tmp):
        from verify.inventory.models import SourceInventory

        out = check_completeness(SourceInventory(file_type="rtf"), extract(make_docx(tmp / "any.docx")))
        assert out.issues == [] and "no completeness comparator" in out.notes[0]

    def test_the_comparison_never_modifies_its_inputs(self, pptx):
        before = pptx.resp.model_dump()
        inv_before = pptx.inv.model_dump()
        pptx.run()
        assert pptx.resp.model_dump() == before and pptx.inv.model_dump() == inv_before

    def test_results_are_deterministic(self, docx, pdf):
        for case in (docx, pdf):
            mutated = drop(case.resp, lambda b: b.type == T.PARAGRAPH and "Revenue" in b.content)
            a, b = case.run(mutated), case.run(mutated)
            assert [m.model_dump() for m in a.matches] == [m.model_dump() for m in b.matches]
            assert [i.model_dump() for i in a.issues] == [i.model_dump() for i in b.issues]

    def test_a_flood_of_missing_objects_is_capped_but_counted(self, tmp):
        import docx as d

        doc = d.Document()
        for i in range(60):
            doc.add_paragraph(f"Distinct paragraph number {i} with its own wording about topic {i * 7}.")
        doc.save(tmp / "many.docx")
        case = make_case(tmp / "many.docx")
        out = case.run(drop(case.resp, lambda b: True))
        miss = [i for i in out.issues if i.code == "possible_missing_text"]
        assert len(miss) == 20 and miss[-1].evidence[0].detail["additional_not_itemised"] == 40


# ====================================================================================== mutation testing


def _single_block_mutations(case: Case, eligible):
    """Remove each eligible block in turn and report which removals Verify missed."""
    base = len(case.run().issues)
    missed = []
    for b in case.resp.blocks:
        if not eligible(b):
            continue
        out = case.run(drop(case.resp, lambda x, bid=b.id: x.id == bid))
        if len(out.issues) <= base:
            missed.append((b.id, b.type.value, b.content[:40]))
    return missed


class TestMutationDetection:
    """Silent-loss rate on injected faults: how many deliberately removed blocks does Verify notice?"""

    def test_xlsx_every_removed_sheet_table_is_noticed(self, xlsx_plain):
        assert _single_block_mutations(xlsx_plain, lambda b: b.type == T.TABLE) == []

    def test_pptx_every_removed_block_is_noticed(self, pptx):
        assert _single_block_mutations(pptx, lambda b: True) == []

    def test_docx_every_removed_block_is_noticed(self, docx):
        assert _single_block_mutations(docx, lambda b: True) == []

    def test_pdf_every_removed_text_table_or_figure_block_on_a_digital_page_is_noticed(self, pdf):
        def eligible(b):
            return b.page in (1, 2) and (b.type in (T.TABLE, T.FIGURE) or (b.type != T.TABLE and len(b.content) >= 15))
        assert _single_block_mutations(pdf, eligible) == []

    def test_detection_rate_summary(self, xlsx_plain, pptx, docx, pdf):
        rates = {}
        for name, case, el in (("xlsx", xlsx_plain, lambda b: b.type == T.TABLE), ("pptx", pptx, lambda b: True),
                               ("docx", docx, lambda b: True),
                               ("pdf", pdf, lambda b: b.page in (1, 2) and len(b.content) >= 15)):
            total = sum(1 for b in case.resp.blocks if el(b))
            missed = len(_single_block_mutations(case, el))
            rates[name] = (total - missed, total)
        assert all(found == total for found, total in rates.values()), rates
