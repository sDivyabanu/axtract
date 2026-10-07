"""Exports (CSV / XLSX / DOCX / Markdown) with spreadsheet formula-injection protection."""

from __future__ import annotations

import csv
import io
from typing import Any, Sequence

_DANGEROUS = ("=", "+", "-", "@", "\t", "\r")


def guard(value: Any) -> Any:
    """Neutralise cells that a spreadsheet would execute: a leading = + - @ gets a protective apostrophe."""
    if isinstance(value, str) and value.startswith(_DANGEROUS):
        return "'" + value
    return value


def csv_bytes(columns: Sequence[str], rows: Sequence[Sequence[Any]]) -> bytes:
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow([guard(c) for c in columns])
    for r in rows:
        w.writerow([guard(c) for c in r])
    return buf.getvalue().encode("utf-8-sig")


def xlsx_bytes(columns: Sequence[str], rows: Sequence[Sequence[Any]], title: str = "Export") -> bytes:
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = (title or "Export")[:31].replace("/", "-")
    ws.append([guard(c) for c in columns])
    for r in rows:
        ws.append([guard(c) for c in r])
    for row in ws.iter_rows():
        for cell in row:
            if isinstance(cell.value, str):
                cell.data_type = "s"  # never let openpyxl treat text as a formula
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def seller_docx(title: str, items: Sequence[dict[str, Any]]) -> bytes:
    import docx

    d = docx.Document()
    d.add_heading(title, 0)
    for it in items:
        p = d.add_paragraph()
        p.add_run(f"{it.get('n', '')}. [{str(it.get('severity', '')).upper()}] ").bold = True
        p.add_run(str(it.get("question", "")))
        if it.get("why"):
            d.add_paragraph("Why: " + str(it["why"])).paragraph_format.left_indent = docx.shared.Inches(0.3)
        ev = ", ".join(f"{e.get('filename', '')} p.{e.get('page') or '?'}" for e in it.get("evidence", []))
        if ev:
            d.add_paragraph("Evidence: " + ev).paragraph_format.left_indent = docx.shared.Inches(0.3)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def seller_markdown(title: str, items: Sequence[dict[str, Any]]) -> bytes:
    lines = [f"# {title}", ""]
    for it in items:
        lines.append(f"{it.get('n', '')}. **[{str(it.get('severity', '')).upper()}]** {it.get('question', '')}")
        if it.get("why"):
            lines.append(f"   - Why: {it['why']}")
        ev = ", ".join(f"{e.get('filename', '')} p.{e.get('page') or '?'}" for e in it.get("evidence", []))
        if ev:
            lines.append(f"   - Evidence: {ev}")
    return ("\n".join(lines) + "\n").encode("utf-8")
