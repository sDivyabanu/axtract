# SECURITY/tests/test_office_scan.py

import os
import zipfile
import tempfile

import openpyxl
import pytest

from SECURITY.office_scan import (
    scan_office,
    scan_office_zip,
    scan_xlsx_formulas,
)


# ============================================================================
# HELPERS
# ============================================================================

def make_xlsx(workbook):
    """Save an openpyxl workbook to a temporary XLSX file."""
    fd, path = tempfile.mkstemp(suffix=".xlsx")
    os.close(fd)
    workbook.save(path)
    return path


def make_fake_office_zip(parts):
    """
    Create a minimal ZIP that looks like an Office Open XML container.

    parts = {
        "xl/vbaProject.bin": b"...",
        "word/document.xml": b"...",
    }
    """
    fd, path = tempfile.mkstemp(suffix=".docx")
    os.close(fd)

    with zipfile.ZipFile(path, "w") as zf:
        for name, content in parts.items():
            zf.writestr(name, content)

    return path


def finding_types(result):
    return [f["type"] for f in result["findings"]]


# ============================================================================
# VBA MACRO DETECTION
# ============================================================================

def test_vba_macro_detected():
    path = make_fake_office_zip({
        "xl/vbaProject.bin": b"fake VBA project"
    })

    try:
        result = scan_office_zip(path)

        assert "office_macro" in finding_types(result)

        finding = next(
            f for f in result["findings"]
            if f["type"] == "office_macro"
        )

        assert finding["severity"] == "high"
        assert finding["action_taken"] == "not_executed"

    finally:
        os.remove(path)


def test_vba_macro_not_executed():
    """
    Security requirement:
    detecting a macro must never execute it.
    """
    path = make_fake_office_zip({
        "xl/vbaProject.bin": b"Sub AutoOpen()\nMsgBox \"HACK\"\nEnd Sub"
    })

    try:
        result = scan_office(path)

        assert "office_macro" in finding_types(result)

        # Scanner should only report it.
        assert all(
            f["action_taken"] == "not_executed"
            for f in result["findings"]
        )

    finally:
        os.remove(path)


# ============================================================================
# OLE DETECTION
# ============================================================================

def test_ole_object_detected():
    path = make_fake_office_zip({
        "word/embeddings/oleObject1.bin": b"fake OLE object"
    })

    try:
        result = scan_office_zip(path)

        assert "office_ole_object" in finding_types(result)

        finding = next(
            f for f in result["findings"]
            if f["type"] == "office_ole_object"
        )

        assert finding["severity"] == "medium"

    finally:
        os.remove(path)


def test_ole_bin_detected():
    path = make_fake_office_zip({
        "word/embeddings/object1.bin": b"fake embedded object"
    })

    try:
        result = scan_office_zip(path)

        assert "office_ole_object" in finding_types(result)

    finally:
        os.remove(path)


# ============================================================================
# DDE DETECTION
# ============================================================================

def test_dde_field_detected():
    xml = b"""
    <w:document>
        <w:p>
            <w:fldSimple>
                <w:instrText>DDEAUTO cmd.exe /c something</w:instrText>
            </w:fldSimple>
        </w:p>
    </w:document>
    """

    path = make_fake_office_zip({
        "word/document.xml": xml
    })

    try:
        result = scan_office_zip(path)

        assert "office_dde_field" in finding_types(result)

        finding = next(
            f for f in result["findings"]
            if f["type"] == "office_dde_field"
        )

        assert finding["severity"] == "high"

    finally:
        os.remove(path)


def test_dde_case_insensitive():
    xml = b"""
    <w:document>
        <w:instrText>ddeauto something</w:instrText>
    </w:document>
    """

    path = make_fake_office_zip({
        "word/document.xml": xml
    })

    try:
        result = scan_office_zip(path)

        assert "office_dde_field" in finding_types(result)

    finally:
        os.remove(path)


# ============================================================================
# CLEAN OFFICE FILE
# ============================================================================

def test_clean_office_zip():
    path = make_fake_office_zip({
        "word/document.xml": b"""
            <w:document>
                <w:p>
                    Revenue increased by 12 percent.
                </w:p>
            </w:document>
        """
    })

    try:
        result = scan_office_zip(path)

        assert result["findings"] == []

    finally:
        os.remove(path)


# ============================================================================
# CORRUPTED / INVALID OFFICE FILE
# ============================================================================

def test_corrupted_office_file():
    fd, path = tempfile.mkstemp(suffix=".docx")
    os.close(fd)

    try:
        with open(path, "wb") as f:
            f.write(b"THIS IS NOT A ZIP FILE")

        result = scan_office_zip(path)

        assert result["findings"] == []
        assert "error" in result
        assert result["error"] is not None

    finally:
        os.remove(path)


# ============================================================================
# F28 — BASIC FORMULA DETECTION
# ============================================================================

def test_formula_is_not_manual_override():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Financials"

    ws["A1"] = "Revenue"
    ws["B1"] = 500

    ws["A2"] = "EBITDA"
    ws["B2"] = "=B1*0.20"

    path = make_xlsx(wb)

    try:
        result = scan_xlsx_formulas(path)

        locations = [
            x["location"]
            for x in result["manual_overrides"]
        ]

        # Revenue is hardcoded and should be detected.
        assert "Financials!B1" in locations

        # EBITDA is formula-backed and should NOT be detected.
        assert "Financials!B2" not in locations

    finally:
        os.remove(path)


# ============================================================================
# F28 — HARD-CODED FINANCIAL VALUE
# ============================================================================

def test_hardcoded_ebitda_detected():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Financials"

    ws["A1"] = "Revenue"
    ws["B1"] = 1000

    ws["A2"] = "EBITDA"
    ws["B2"] = 250

    path = make_xlsx(wb)

    try:
        result = scan_xlsx_formulas(path)

        assert "Financials!B2" in [
            x["location"]
            for x in result["manual_overrides"]
        ]

        finding = next(
            f for f in result["findings"]
            if f["block_id"] == "Financials!B2"
        )

        assert finding["type"] == "manual_override_suspected"
        assert finding["severity"] == "high"

    finally:
        os.remove(path)


# ============================================================================
# F28 — REALISTIC PE / IB FINANCIAL MODEL
# ============================================================================

def test_pe_financial_model():
    """
    Synthetic PE-style financial model.

    Formula-backed metrics should be accepted.
    Hardcoded financial values should be surfaced.
    """

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "LBO Model"

    ws["A1"] = "Metric"
    ws["B1"] = "FY2024"
    ws["C1"] = "FY2025"

    ws["A2"] = "Revenue"
    ws["B2"] = 850
    ws["C2"] = 1020

    ws["A3"] = "EBITDA"
    ws["B3"] = 170

    # Deliberately hardcoded instead of formula.
    ws["C3"] = 999

    ws["A4"] = "EBITDA Margin"
    ws["B4"] = "=B3/B2"
    ws["C4"] = "=C3/C2"

    ws["A5"] = "Net Debt"
    ws["B5"] = 500
    ws["C5"] = 450

    path = make_xlsx(wb)

    try:
        result = scan_xlsx_formulas(path)

        locations = [
            x["location"]
            for x in result["manual_overrides"]
        ]

        # Deliberately suspicious EBITDA.
        assert "LBO Model!C3" in locations

        # Formula-backed margins should not be flagged.
        assert "LBO Model!B4" not in locations
        assert "LBO Model!C4" not in locations

    finally:
        os.remove(path)


# ============================================================================
# F28 — CLEAN FORMULA MODEL
# ============================================================================

def test_clean_formula_model():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Model"

    ws["A1"] = "Revenue"
    ws["B1"] = 1000

    ws["A2"] = "COGS"
    ws["B2"] = 600

    ws["A3"] = "Gross Profit"
    ws["B3"] = "=B1-B2"

    ws["A4"] = "EBITDA"
    ws["B4"] = "=B3*0.30"

    path = make_xlsx(wb)

    try:
        result = scan_xlsx_formulas(path)

        # Formula-backed financial metrics should not be flagged.
        assert "Model!B3" not in [
            x["location"] for x in result["manual_overrides"]
        ]

        assert "Model!B4" not in [
            x["location"] for x in result["manual_overrides"]
        ]

    finally:
        os.remove(path)


# ============================================================================
# F28 — NON-FINANCIAL HARD-CODED VALUES
# ============================================================================

def test_non_financial_hardcode_not_flagged():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"

    ws["A1"] = "Employee ID"
    ws["B1"] = 10025

    ws["A2"] = "Department Code"
    ws["B2"] = 42

    path = make_xlsx(wb)

    try:
        result = scan_xlsx_formulas(path)

        assert result["manual_overrides"] == []

    finally:
        os.remove(path)


# ============================================================================
# F28 — MULTIPLE SHEETS
# ============================================================================

def test_multiple_sheets_scanned():
    wb = openpyxl.Workbook()

    ws1 = wb.active
    ws1.title = "Income Statement"

    ws1["A1"] = "Revenue"
    ws1["B1"] = 1000

    ws2 = wb.create_sheet("Balance Sheet")

    ws2["A1"] = "Assets"
    ws2["B1"] = 5000

    path = make_xlsx(wb)

    try:
        result = scan_xlsx_formulas(path)

        locations = [
            x["location"]
            for x in result["manual_overrides"]
        ]

        assert "Income Statement!B1" in locations
        assert "Balance Sheet!B1" in locations

    finally:
        os.remove(path)


# ============================================================================
# F28 — CLEAN XLSX
# ============================================================================

def test_clean_xlsx():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Financials"

    ws["A1"] = "Revenue"
    ws["B1"] = "=1000"

    ws["A2"] = "EBITDA"
    ws["B2"] = "=B1*0.2"

    path = make_xlsx(wb)

    try:
        result = scan_office(path)

        assert result["clean"] is True
        assert result["findings"] == []
        assert result["manual_overrides"] == []

    finally:
        os.remove(path)


# ============================================================================
# COMBINED SECURITY TEST
# ============================================================================

def test_macro_and_ole_both_detected():
    path = make_fake_office_zip({
        "xl/vbaProject.bin": b"fake macro",
        "xl/embeddings/oleObject1.bin": b"fake OLE",
    })

    try:
        result = scan_office(path)

        types = finding_types(result)

        assert "office_macro" in types
        assert "office_ole_object" in types

    finally:
        os.remove(path)


# ============================================================================
# RESULT STRUCTURE
# ============================================================================

def test_result_schema():
    wb = openpyxl.Workbook()
    ws = wb.active

    ws["A1"] = "Revenue"
    ws["B1"] = 500

    path = make_xlsx(wb)

    try:
        result = scan_office(path)

        assert "findings" in result
        assert "manual_overrides" in result
        assert "clean" in result

        assert isinstance(result["findings"], list)
        assert isinstance(result["manual_overrides"], list)
        assert isinstance(result["clean"], bool)

    finally:
        os.remove(path)